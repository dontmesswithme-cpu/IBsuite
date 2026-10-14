from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter, range_boundaries
from PIL import Image as PILImage


def _parse_range(value: str, what: str) -> tuple[int, int, int, int]:
    text = (value or "").strip().replace("$", "")
    if not text:
        raise ValueError(f"bad {what}: {value!r}")
    try:
        return range_boundaries(text)
    except Exception as exc:
        raise ValueError(f"bad {what}: {value!r}") from exc


def _parse_anchor(value: str, what: str) -> str:
    bounds = _parse_range(value, f"{what} anchor")
    if bounds[0] != bounds[2] or bounds[1] != bounds[3]:
        raise ValueError(f"invalid {what} anchor: {value!r} (expected a single cell)")
    return f"{get_column_letter(bounds[0])}{bounds[1]}"


def _range_ref(bounds: tuple[int, int, int, int]) -> str:
    min_col, min_row, max_col, max_row = bounds
    start = f"{get_column_letter(min_col)}{min_row}"
    end = f"{get_column_letter(max_col)}{max_row}"
    return start if start == end else f"{start}:{end}"


def _open_wb(path: str) -> Workbook:
    target = Path(path)
    if not target.is_file():
        raise FileNotFoundError(f"workbook not found: {target}")
    return load_workbook(target)


def _worksheet(wb: Workbook, sheet: str):
    if sheet not in wb.sheetnames:
        raise ValueError(f"unknown sheet: {sheet!r}")
    return wb[sheet]


def _make_chart(chart_type: str):
    key = (chart_type or "").strip().lower()
    if key == "line":
        return LineChart()
    if key == "bar":
        chart = BarChart()
        chart.type = "bar"
        return chart
    if key == "column":
        chart = BarChart()
        chart.type = "col"
        return chart
    if key == "pie":
        return PieChart()
    if key == "area":
        return AreaChart()
    raise ValueError(
        f"unknown chart type: {chart_type!r} (expected one of line, bar, column, pie, area)"
    )


def _sheet_token(name: str) -> str:
    if name and all(ch.isalnum() or ch == "_" for ch in name):
        return name
    return f"'{name.replace(chr(39), chr(39) * 2)}'"


def add_chart(path: str, sheet: str, data_range: str, chart_type: str, anchor: str) -> dict:
    bounds = _parse_range(data_range, "data range")
    cell = _parse_anchor(anchor, "chart")
    chart = _make_chart(chart_type)

    min_col, min_row, max_col, max_row = bounds
    wb = _open_wb(path)
    try:
        ws = _worksheet(wb, sheet)
        if max_col - min_col >= 1:
            data = Reference(ws, min_col=min_col + 1, min_row=min_row, max_col=max_col, max_row=max_row)
            cats = Reference(ws, min_col=min_col, min_row=min_row + 1, max_row=max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
        else:
            data = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
            chart.add_data(data, titles_from_data=True)
        chart.title = f"{sheet} {chart_type}"
        ws.add_chart(chart, cell)
        wb.save(path)
        return {"ok": True, "chart": chart_type, "sheet": sheet, "ref": _range_ref(bounds), "anchor": cell}
    finally:
        wb.close()


def add_image(path: str, sheet: str, image_path: str, anchor: str) -> dict:
    cell = _parse_anchor(anchor, "image")
    source = Path(image_path)
    if not source.is_file():
        raise FileNotFoundError(f"image not found: {source}")
    try:
        with PILImage.open(source) as probe:
            probe.verify()
    except Exception as exc:
        raise ValueError(f"unreadable image file: {source}") from exc

    wb = _open_wb(path)
    try:
        ws = _worksheet(wb, sheet)
        ws.add_image(XLImage(str(source)), cell)
        wb.save(path)
        return {"ok": True, "sheet": sheet, "image": str(source), "anchor": cell}
    finally:
        wb.close()


def add_pivot(
    path: str,
    sheet: str,
    source_range: str,
    dest: str,
    rows: list[str],
    values: list[str],
) -> dict:
    if not rows:
        raise ValueError("pivot requires at least one row field in 'rows'")
    if not values:
        raise ValueError("pivot requires at least one value field in 'values'")

    bounds = _parse_range(source_range, "source range")

    if "!" in dest:
        dest_sheet, _, dest_cell = dest.partition("!")
    else:
        dest_sheet, dest_cell = dest, "A1"
    dest_sheet = dest_sheet.strip().strip("'\"")
    if not dest_sheet:
        raise ValueError(f"invalid pivot destination: {dest!r}")
    dest_bounds = _parse_range(dest_cell.strip(), "destination")
    if dest_bounds[0] != dest_bounds[2] or dest_bounds[1] != dest_bounds[3]:
        raise ValueError(f"invalid pivot destination: {dest!r} (expected Sheet!Cell)")

    min_col, min_row, max_col, max_row = bounds
    if max_row - min_row < 1:
        raise ValueError(f"pivot source needs a header row and data: {source_range!r}")

    wb = _open_wb(path)
    try:
        ws = _worksheet(wb, sheet)
        headers = [ws.cell(row=min_row, column=c).value for c in range(min_col, max_col + 1)]

        if rows[0] not in headers:
            raise ValueError(f"unknown pivot row header: {rows[0]!r}")
        missing = [v for v in values if v not in headers]
        if missing:
            raise ValueError(f"unknown pivot value headers: {missing!r}")

        group_col = min_col + headers.index(rows[0])
        value_cols = [min_col + headers.index(v) for v in values]

        keys: list = []
        for r in range(min_row + 1, max_row + 1):
            key = ws.cell(row=r, column=group_col).value
            if key not in keys:
                keys.append(key)

        dws = wb[dest_sheet] if dest_sheet in wb.sheetnames else None
        if dws is not None:
            raise ValueError(
                f"pivot destination sheet already exists: {dest_sheet!r} "
                "(choose a new sheet name)"
            )
        dws = wb.create_sheet(dest_sheet)
        start_col, start_row = dest_bounds[0], dest_bounds[1]
        token = _sheet_token(sheet)
        value_range = lambda c: f"{token}!${get_column_letter(c)}${min_row + 1}:${get_column_letter(c)}${max_row}"
        group_range = value_range(group_col)
        key_col_letter = get_column_letter(start_col)

        dws.cell(row=start_row, column=start_col, value=rows[0])
        for offset, header in enumerate(values):
            dws.cell(row=start_row, column=start_col + 1 + offset, value=f"SUM of {header}")

        for index, key in enumerate(keys):
            row = start_row + 1 + index
            dws.cell(row=row, column=start_col, value=key)
            for offset, vcol in enumerate(value_cols):
                letter = get_column_letter(vcol)
                dws.cell(
                    row=row,
                    column=start_col + 1 + offset,
                    value=f"=SUMIFS({value_range(vcol)},{group_range},${key_col_letter}${row})",
                )

        wb.save(path)
        return {
            "ok": True,
            "sheet": dest_sheet,
            "cell": f"{key_col_letter}{start_row}",
            "warning": "native Excel pivot tables cannot be authored headless; wrote a SUMIFS summary sheet instead",
        }
    finally:
        wb.close()
