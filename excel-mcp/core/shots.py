import base64
import io
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from PIL import Image, ImageDraw, ImageFont

from config import resolve_style

THUMB_WIDTH = 320
THUMB_MAX_HEIGHT = 480
FULL_WIDTH = 1600
MAX_THUMB_B64 = 40000
MAX_COLS = 50
MAX_ROWS = 200
DEFAULT_COL_WIDTH = 8.43
DEFAULT_ROW_HEIGHT = 15.0


def _load_font(size, bold: bool):
    size = max(8, int(size or 10))
    candidates = []
    if bold:
        candidates.append(r"C:\Windows\Fonts\arialbd.ttf")
    candidates.append(r"C:\Windows\Fonts\arial.ttf")
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            continue
    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()


def _hex6(value, default: str) -> str:
    if not value or value == "00000000":
        return default
    text = str(value).upper()
    return text[-6:] if len(text) >= 6 else default


def _font_hex(color_obj) -> str:
    try:
        if getattr(color_obj, "type", None) != "rgb":
            return "000000"
        rgb = color_obj.rgb
    except AttributeError:
        return "000000"
    return _hex6(rgb, "000000")


def _fill_hex(cell) -> str:
    if cell.fill.patternType != "solid":
        return "FFFFFF"
    try:
        fg = cell.fill.fgColor
        if getattr(fg, "type", None) != "rgb":
            return "FFFFFF"
        rgb = fg.rgb
    except AttributeError:
        return "FFFFFF"
    return _hex6(rgb, "FFFFFF")


def _display(value, number_format) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return f"{value:,}"
    if isinstance(value, float):
        if value.is_integer():
            return f"{int(value):,}"
        fmt = str(number_format or "")
        if "%" in fmt:
            return f"{value:.1%}"
        if "0.0" in fmt:
            return f"{value:,.1f}"
        return f"{value:,.2f}"
    text = str(value)
    return text if len(text) <= 60 else text[:57] + "..."


def _parse_range(value: str) -> tuple[int, int, int, int]:
    text = (value or "").strip().replace("$", "")
    if not text:
        raise ValueError(f"bad cell range: {value!r}")
    try:
        return range_boundaries(text)
    except Exception as exc:
        raise ValueError(f"bad cell range: {value!r}") from exc


def _render(path: str, sheet: str, cell_range: str | None) -> Image.Image:
    target = Path(path)
    if not target.is_file():
        raise FileNotFoundError(f"workbook not found: {target}")

    if cell_range is None:
        bounds = None
    else:
        bounds = _parse_range(cell_range)

    wb = load_workbook(path)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"unknown sheet: {sheet!r}")
        ws = wb[sheet]

        if bounds is None:
            bounds = range_boundaries(ws.dimensions)
        min_col, min_row, max_col, max_row = bounds
        max_col = min(max_col, min_col + MAX_COLS - 1)
        max_row = min(max_row, min_row + MAX_ROWS - 1)

        col_px = []
        for col in range(min_col, max_col + 1):
            width = ws.column_dimensions[get_column_letter(col)].width
            if width is None:
                width = DEFAULT_COL_WIDTH
            col_px.append(max(30, int(width * 7 + 5)))

        row_px = []
        for row in range(min_row, max_row + 1):
            height = ws.row_dimensions[row].height
            if height is None:
                height = DEFAULT_ROW_HEIGHT
            row_px.append(max(14, int(height * 4 / 3)))

        grid_hex = resolve_style("border_grid", None)
        grid = "#" + (grid_hex if isinstance(grid_hex, str) else "D9D9D9")

        img = Image.new("RGB", (sum(col_px), sum(row_px)), "#FFFFFF")
        draw = ImageDraw.Draw(img)

        y = 0
        for row_index, row in enumerate(range(min_row, max_row + 1)):
            height = row_px[row_index]
            x = 0
            for col_index, col in enumerate(range(min_col, max_col + 1)):
                width = col_px[col_index]
                cell = ws.cell(row=row, column=col)

                fill = _fill_hex(cell)
                if fill != "FFFFFF":
                    draw.rectangle([x, y, x + width - 1, y + height - 1], fill="#" + fill)
                draw.rectangle([x, y, x + width - 1, y + height - 1], outline=grid)

                text = _display(cell.value, cell.number_format)
                if text:
                    font = _load_font(cell.font.size, bool(cell.font.bold))
                    halign = (cell.alignment.horizontal or "").lower()
                    if not halign:
                        halign = (
                            "right"
                            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
                            else "left"
                        )
                    text_width = draw.textlength(text, font=font)
                    if halign == "right":
                        tx = x + width - 3 - text_width
                    elif halign == "center":
                        tx = x + (width - text_width) / 2
                    else:
                        tx = x + 3
                    draw.text((tx, y + 2), text, font=font, fill="#" + _font_hex(cell.font.color))
                x += width
            y += height
        return img
    finally:
        wb.close()


def _to_b64(img: Image.Image) -> str:
    buffer = io.BytesIO()
    img.save(buffer, format="PNG", optimize=True)
    return base64.b64encode(buffer.getvalue()).decode("ascii")


def _fit(img: Image.Image, width: int) -> Image.Image:
    current, height = img.size
    if current == width:
        return img
    return img.resize((width, max(1, int(height * width / current))), Image.LANCZOS)


def _headed_copy(path: str, sheet: str, cell_range: str | None):
    try:
        import win32com.client
        from PIL import ImageGrab
    except ImportError:
        return None
    target = Path(path)
    if not target.is_file():
        return None
    try:
        app = win32com.client.Dispatch("Excel.Application")
    except Exception:
        return None
    try:
        app.Visible = False
        app.DisplayAlerts = False
        wb = app.Workbooks.Open(str(target.resolve()), ReadOnly=True)
        try:
            try:
                ws = wb.Worksheets(sheet)
            except Exception:
                return None
            if cell_range:
                try:
                    bounds = range_boundaries(cell_range.strip().replace("$", ""))
                except Exception:
                    return None
                a = f"{get_column_letter(bounds[0])}{bounds[1]}"
                b = f"{get_column_letter(bounds[2])}{bounds[3]}"
                ref = a if a == b else f"{a}:{b}"
            else:
                ref = ws.UsedRange.Address
            ws.Range(ref).CopyPicture(Appearance=1, Format=2)
            shot = ImageGrab.grabclipboard()
            if shot is None:
                return None
            return shot.convert("RGB")
        finally:
            try:
                wb.Close(SaveChanges=False)
            except Exception:
                pass
    except Exception:
        return None
    finally:
        try:
            app.Quit()
        except Exception:
            pass


def thumbnail(path: str, sheet: str, cell_range: str | None) -> str:
    img = _render(path, sheet, cell_range)
    small = _fit(img, THUMB_WIDTH)
    if small.size[1] > THUMB_MAX_HEIGHT:
        small = small.crop((0, 0, small.size[0], THUMB_MAX_HEIGHT))
    small = small.convert("P", palette=Image.ADAPTIVE, colors=32).convert("RGB")
    out = _to_b64(small)
    while len(out) >= MAX_THUMB_B64 and small.size[0] > 160:
        small = small.resize((small.size[0] // 2, small.size[1] // 2), Image.LANCZOS)
        small = small.convert("P", palette=Image.ADAPTIVE, colors=16).convert("RGB")
        out = _to_b64(small)
    return out


def full_shot(path: str, sheet: str, cell_range: str | None, mode: str = "headless") -> str:
    if mode not in ("headless", "headed"):
        raise ValueError(f"unknown mode: {mode!r} (expected 'headless' or 'headed')")
    if mode == "headed":
        shot = _headed_copy(path, sheet, cell_range)
        if shot is not None:
            return _to_b64(_fit(shot, FULL_WIDTH))
    return _to_b64(_fit(_render(path, sheet, cell_range), FULL_WIDTH))
