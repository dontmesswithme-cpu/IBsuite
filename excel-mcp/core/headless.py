from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

MAX_CELLS = 100_000


def _bounds(cell_range: str) -> tuple[int, int, int, int]:
    text = (cell_range or "").strip().replace("$", "")
    if not text:
        raise ValueError(f"bad cell range: {cell_range!r}")
    try:
        return range_boundaries(text)
    except Exception as exc:
        raise ValueError(f"bad cell range: {cell_range!r}") from exc


def _single_cell_ref(bounds: tuple[int, int, int, int]) -> str:
    min_col, min_row, _, _ = bounds
    return f"{get_column_letter(min_col)}{min_row}"


def _open_existing(path: str) -> Workbook:
    target = Path(path)
    if not target.is_file():
        raise FileNotFoundError(f"workbook not found: {target}")
    return load_workbook(target)


def _sheet(wb: Workbook, sheet: str):
    if sheet not in wb.sheetnames:
        raise ValueError(f"unknown sheet: {sheet!r}")
    return wb[sheet]


def _normalise_values(values) -> list[list]:
    if not isinstance(values, (list, tuple)) or len(values) == 0:
        raise ValueError("values must be a non-empty list of rows")
    rows = []
    for row in values:
        if not isinstance(row, (list, tuple)) or len(row) == 0:
            raise ValueError("values must be a non-empty list of rows")
        rows.append(list(row))
    width = len(rows[0])
    for row in rows:
        if len(row) != width:
            raise ValueError("ragged values: every row must have the same length")
    return rows


def _read_grid(ws, bounds: tuple[int, int, int, int]) -> dict:
    min_col, min_row, max_col, max_row = bounds
    values: list[list] = []
    formulas: list[list] = []
    number_formats: list[list] = []

    for row in range(min_row, max_row + 1):
        value_row: list = []
        formula_row: list = []
        format_row: list = []
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row=row, column=col)
            value = cell.value
            value_row.append(value)
            formula_row.append(value if isinstance(value, str) and value.startswith("=") else None)
            format_row.append(cell.number_format)
        values.append(value_row)
        formulas.append(formula_row)
        number_formats.append(format_row)

    return {"values": values, "formulas": formulas, "number_formats": number_formats}


def _check_shape(bounds: tuple[int, int, int, int], rows: list[list], label: str) -> None:
    min_col, min_row, max_col, max_row = bounds
    height = max_row - min_row + 1
    width = max_col - min_col + 1
    if height == 1 and width == 1:
        return
    if len(rows) != height or len(rows[0]) != width:
        raise ValueError(
            f"{label}: expects {height}x{width} values, got {len(rows)}x{len(rows[0])}"
        )


def _guard_cell_count(rows: list[list]) -> None:
    total = len(rows) * len(rows[0])
    if total > MAX_CELLS:
        raise ValueError(
            f"write limited to {MAX_CELLS} cells per call, split into batches (got {total})"
        )


def list_sheets(path: str) -> list[dict]:
    wb = _open_existing(path)
    try:
        described = []
        for name in wb.sheetnames:
            ws = wb[name]
            described.append(
                {
                    "name": name,
                    "dims": ws.dimensions,
                    "frozen": ws.freeze_panes,
                    "print": {
                        "orientation": ws.page_setup.orientation,
                        "fit_to_pages_wide": ws.page_setup.fitToWidth,
                        "fit_to_pages_tall": ws.page_setup.fitToHeight,
                    },
                }
            )
        return described
    finally:
        wb.close()


def read_range(path: str, sheet: str, cell_range: str) -> dict:
    bounds = _bounds(cell_range)
    wb = _open_existing(path)
    try:
        ws = _sheet(wb, sheet)
        return _read_grid(ws, bounds)
    finally:
        wb.close()


def write_range(path: str, sheet: str, cell_range: str, values: list[list]) -> dict:
    rows = _normalise_values(values)
    bounds = _bounds(cell_range)
    _check_shape(bounds, rows, "write_range")
    _guard_cell_count(rows)

    min_col, min_row, _, _ = bounds
    wb = _open_existing(path)
    try:
        ws = _sheet(wb, sheet)
        for row_offset, row in enumerate(rows):
            for col_offset, value in enumerate(row):
                ws.cell(row=min_row + row_offset, column=min_col + col_offset).value = value
        wb.save(path)
        return {"ok": True, "rows": len(rows), "cols": len(rows[0])}
    finally:
        wb.close()


def set_formula(path: str, sheet: str, cell: str, formula: str) -> dict:
    if not isinstance(formula, str) or not formula.strip().startswith("="):
        raise ValueError(f"formula must start with '=': {formula!r}")

    bounds = _bounds(cell)
    if bounds[0] != bounds[2] or bounds[1] != bounds[3]:
        raise ValueError(f"bad cell range: single cell expected, got {cell!r}")

    ref = _single_cell_ref(bounds)
    wb = _open_existing(path)
    try:
        ws = _sheet(wb, sheet)
        ws[ref].value = formula
        wb.save(path)
        return {"ok": True, "sheet": sheet, "cell": ref, "formula": formula}
    finally:
        wb.close()


def add_sheet(path: str, name: str) -> dict:
    title = (name or "").strip()
    if not title:
        raise ValueError("empty sheet name")

    wb = _open_existing(path)
    try:
        if title in wb.sheetnames:
            raise ValueError(f"sheet already exists: {title!r}")
        wb.create_sheet(title)
        wb.save(path)
        return {"ok": True, "sheet": title}
    finally:
        wb.close()


def _preview_target(ops: dict) -> tuple[tuple[int, int, int, int], list[list]]:
    tool = ops.get("tool", "write_range")

    if tool == "write_range":
        rows = _normalise_values(ops.get("values"))
        bounds = _bounds(ops["cell_range"])
        _check_shape(bounds, rows, "write_range")
        _guard_cell_count(rows)

        min_col, min_row, max_col, max_row = bounds
        if (min_row, min_col) == (max_row, max_col):
            height, width = len(rows), len(rows[0])
        else:
            height = max_row - min_row + 1
            width = max_col - min_col + 1
        return (min_col, min_row, min_col + width - 1, min_row + height - 1), rows

    if tool == "set_formula":
        formula = ops.get("formula")
        if not isinstance(formula, str) or not formula.strip().startswith("="):
            raise ValueError(f"formula must start with '=': {formula!r}")
        bounds = _bounds(ops["cell"])
        if bounds[0] != bounds[2] or bounds[1] != bounds[3]:
            raise ValueError(f"bad cell range: single cell expected, got {ops['cell']!r}")
        return bounds, [[formula]]

    raise ValueError(f"unsupported preview_diff tool: {tool!r}")


def preview_diff(path: str, ops: dict) -> dict:
    for key in ("sheet", "tool"):
        if key not in ops:
            raise ValueError(f"preview_diff requires {key!r} in ops")

    if "cell_range" not in ops and "cell" not in ops:
        raise ValueError("preview_diff requires 'cell_range' (or 'cell') in ops")

    bounds, after_rows = _preview_target(ops)

    wb = _open_existing(path)
    try:
        ws = _sheet(wb, ops["sheet"])
        before = _read_grid(ws, bounds)
    finally:
        wb.close()

    diff = []
    for row_offset, row in enumerate(after_rows):
        for col_offset, new_value in enumerate(row):
            min_col, min_row = bounds[0], bounds[1]
            ref = f"{get_column_letter(min_col + col_offset)}{min_row + row_offset}"
            old_value = before["values"][row_offset][col_offset]
            if old_value != new_value:
                diff.append({"cell": ref, "before": old_value, "after": new_value})

    return {
        "tool": ops.get("tool", "write_range"),
        "sheet": ops["sheet"],
        "diff": diff,
        "before": before,
        "after": {"values": after_rows},
    }
