from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

from config import resolve_style

STYLE_KEYS = (
    "font",
    "fill",
    "border",
    "number_format",
    "alignment",
    "col_widths",
    "row_height",
    "freeze",
    "print_fit",
    "table_style",
    "validation",
    "conditional",
)

_COLOR_NAMES = {
    "white": "FFFFFF",
    "black": "000000",
    "grey": "808080",
    "gray": "808080",
    "blue": "0000FF",
    "red": "FF0000",
    "navy": "002060",
}

_FONT_FAMILIES = (
    "Arial",
    "Calibri",
    "Cambria",
    "Consolas",
    "Courier New",
    "Georgia",
    "Helvetica",
    "Segoe UI",
    "Tahoma",
    "Times New Roman",
    "Verdana",
)
_FONT_ALIAS = {name.lower(): name for name in _FONT_FAMILIES}

_FONT_FLAGS = {"bold", "b", "italic", "i", "underline", "u"}

_BORDER_STYLES = (
    "dashDot",
    "dashDotDot",
    "dashed",
    "dotted",
    "double",
    "hair",
    "medium",
    "mediumDashDot",
    "mediumDashDotDot",
    "mediumDashed",
    "slantDashDot",
    "thick",
    "thin",
)

_ALIGNMENT_KEYS = (
    "horizontal",
    "vertical",
    "wrap_text",
    "shrink_to_fit",
    "indent",
    "text_rotation",
)


def _parse_range(cell_range: str) -> tuple[int, int, int, int]:
    text = (cell_range or "").strip().replace("$", "")
    if not text:
        raise ValueError(f"bad cell range: {cell_range!r}")
    try:
        return range_boundaries(text)
    except Exception as exc:
        raise ValueError(f"bad cell range: {cell_range!r}") from exc


def _range_ref(bounds: tuple[int, int, int, int]) -> str:
    min_col, min_row, max_col, max_row = bounds
    start = f"{get_column_letter(min_col)}{min_row}"
    end = f"{get_column_letter(max_col)}{max_row}"
    return start if start == end else f"{start}:{end}"


def _open_wb(path: str) -> Workbook:
    target = Path(path)
    if not target.is_file():
        raise FileNotFoundError(f"workbook not found: {target}")
    return load_workbook(target)


def _worksheet(wb: Workbook, sheet: str):
    if sheet not in wb.sheetnames:
        raise ValueError(f"unknown sheet: {sheet!r}")
    return wb[sheet]


def _colour_code(value, what: str) -> str | None:
    if value is None:
        return None
    if isinstance(value, str) and not value.strip():
        return None
    text = str(value).strip()
    named = _COLOR_NAMES.get(text.lower())
    if named is not None:
        return named
    body = text[1:] if text.startswith("#") else text
    if len(body) == 8:
        body = body[2:]
    if len(body) != 6 or any(c not in "0123456789ABCDEFabcdef" for c in body):
        raise ValueError(f"invalid {what}: {value!r}")
    return body.upper()


def _positive_number(value, what: str) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"invalid {what}: {value!r}") from exc
    if number <= 0:
        raise ValueError(f"invalid {what}: {value!r} (must be positive)")
    return number


def _parse_font(spec) -> dict:
    if spec is None:
        return {"name": None, "size": None, "bold": False, "italic": False, "underline": False, "color": None}
    if isinstance(spec, dict):
        out = {"name": None, "size": None, "bold": False, "italic": False, "underline": False, "color": None}
        if spec.get("name") is not None:
            out["name"] = str(spec["name"])
        if spec.get("size") is not None:
            out["size"] = _positive_number(spec["size"], "font size")
        for flag in ("bold", "italic", "underline"):
            if flag in spec:
                out[flag] = bool(spec[flag])
        if spec.get("color") is not None:
            out["color"] = _colour_code(spec["color"], "font colour")
        return out

    tokens = str(spec).split()
    if not tokens:
        raise ValueError(f"invalid font spec: {spec!r}")

    out = {"name": None, "size": None, "bold": False, "italic": False, "underline": False, "color": None}
    name_parts: list[str] = []
    for token in tokens:
        low = token.lower()
        if low in ("bold", "b"):
            out["bold"] = True
        elif low in ("italic", "i"):
            out["italic"] = True
        elif low in ("underline", "u"):
            out["underline"] = True
        elif low in _COLOR_NAMES:
            out["color"] = _COLOR_NAMES[low]
        else:
            try:
                out["size"] = _positive_number(token, "font size")
            except ValueError:
                name_parts.append(token)

    if name_parts:
        family = " ".join(name_parts)
        canonical = _FONT_ALIAS.get(family.lower())
        if canonical is None:
            raise ValueError(f"invalid font spec: {spec!r} (unknown family {family!r})")
        out["name"] = canonical
    return out


def _merge_font(base: dict, override: dict) -> dict:
    merged = dict(base)
    for key in ("name", "size", "color"):
        if override.get(key) is not None:
            merged[key] = override[key]
    for flag in ("bold", "italic", "underline"):
        if override.get(flag):
            merged[flag] = True
    return merged


def _to_font(parsed: dict) -> Font:
    kwargs: dict = {}
    if parsed["name"] is not None:
        kwargs["name"] = parsed["name"]
    if parsed["size"] is not None:
        kwargs["size"] = parsed["size"]
    kwargs["bold"] = parsed["bold"]
    kwargs["italic"] = parsed["italic"]
    kwargs["underline"] = "single" if parsed["underline"] else None
    if parsed["color"] is not None:
        kwargs["color"] = parsed["color"]
    return Font(**kwargs)


def _parse_border(spec) -> Border:
    grid_color = resolve_style("border_grid", None)
    grid_style = resolve_style("border_grid_style", None)

    style_name = grid_style
    colour = grid_color
    if isinstance(spec, str):
        style_name = spec
    elif isinstance(spec, dict):
        if "color" in spec and spec["color"] is not None:
            colour = _colour_code(spec["color"], "border colour")
        if "style" in spec and spec["style"] is not None:
            style_name = spec["style"]
    elif spec is not None:
        raise ValueError(f"invalid border spec: {spec!r}")

    if style_name not in _BORDER_STYLES:
        raise ValueError(f"invalid border style: {style_name!r}")
    colour = _colour_code(colour, "border colour")

    side = Side(style=style_name, color=colour)
    return Border(left=side, right=side, top=side, bottom=side)


def _parse_alignment(spec) -> Alignment:
    if not isinstance(spec, dict):
        raise ValueError(f"invalid alignment spec: {spec!r}")
    unknown = [key for key in spec if key not in _ALIGNMENT_KEYS]
    if unknown:
        raise ValueError(f"invalid alignment keys: {unknown!r}")
    return Alignment(**spec)


def _paint_cell(cell, *, font: Font | None, fill: PatternFill | None, border: Border | None, number_format: str | None) -> None:
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if border is not None:
        cell.border = border
    if number_format is not None:
        cell.number_format = number_format


def _paint_range(ws, bounds, *, font=None, fill=None, border=None, number_format=None, alignment=None) -> None:
    min_col, min_row, max_col, max_row = bounds
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row=row, column=col)
            _paint_cell(cell, font=font, fill=fill, border=border, number_format=number_format)
            if alignment is not None:
                cell.alignment = alignment


def _set_dimensions(ws, bounds, *, col_widths, row_height) -> None:
    min_col, min_row, max_col, max_row = bounds
    if col_widths is not None:
        if isinstance(col_widths, dict):
            for key, raw_width in col_widths.items():
                width = _positive_number(raw_width, "column width")
                if isinstance(key, int):
                    letter = get_column_letter(key)
                else:
                    letter = str(key).strip().upper().replace("$", "")
                    if not letter.isalpha() or len(letter) > 3:
                        raise ValueError(f"invalid column key: {key!r}")
                ws.column_dimensions[letter].width = width
        else:
            if isinstance(col_widths, (int, float)):
                widths = [_positive_number(col_widths, "column width")] * (max_col - min_col + 1)
            else:
                widths = [_positive_number(w, "column width") for w in col_widths]
                if len(widths) != max_col - min_col + 1:
                    raise ValueError(
                        f"col_widths expects {max_col - min_col + 1} entries, got {len(widths)}"
                    )
            for offset, width in enumerate(widths):
                ws.column_dimensions[get_column_letter(min_col + offset)].width = width
    if row_height is not None:
        height = _positive_number(row_height, "row height")
        for row in range(min_row, max_row + 1):
            ws.row_dimensions[row].height = height


def _apply_print(ws, spec: dict | None) -> None:
    if spec is None:
        return
    if not isinstance(spec, dict):
        raise ValueError(f"invalid print spec: {spec!r}")
    if "orientation" in spec:
        if spec["orientation"] not in ("landscape", "portrait"):
            raise ValueError(f"invalid print orientation: {spec['orientation']!r}")
        ws.page_setup.orientation = spec["orientation"]
    if "fit_to_pages_wide" in spec:
        ws.page_setup.fitToWidth = int(spec["fit_to_pages_wide"])
    if "fit_to_pages_tall" in spec:
        ws.page_setup.fitToHeight = int(spec["fit_to_pages_tall"])
    if "footer" in spec:
        ws.oddFooter.left.text = None
        ws.oddFooter.center.text = str(spec["footer"])
        ws.oddFooter.right.text = None
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def _add_table(ws, ref: str, style_name: str) -> str:
    if not isinstance(style_name, str) or not style_name.strip():
        raise ValueError(f"invalid table style: {style_name!r}")
    base = "".join(ch for ch in f"IB{ws.title}{ref}" if ch.isalnum()) or "IBTable"
    name = base[:60]
    suffix = 1
    taken = set(ws.tables)
    while name in taken:
        suffix += 1
        name = f"{base[:55]}{suffix}"
    ws.add_table(
        Table(
            displayName=name,
            ref=ref,
            tableStyleInfo=TableStyleInfo(
                name=style_name.strip(),
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=False,
                showColumnStripes=False,
            ),
        )
    )
    return name


def _add_validation(ws, ref: str, spec: dict) -> None:
    if not isinstance(spec, dict):
        raise ValueError(f"invalid validation spec: {spec!r}")
    allowed = {"type", "formula1", "formula2", "allow_blank", "show_error", "error", "error_title", "prompt", "prompt_title", "operator"}
    unknown = [key for key in spec if key not in allowed]
    if unknown:
        raise ValueError(f"invalid validation keys: {unknown!r}")
    dv = DataValidation(
        type=spec.get("type", "list"),
        formula1=spec.get("formula1"),
        formula2=spec.get("formula2"),
        allow_blank=spec.get("allow_blank", True),
        showErrorMessage=spec.get("show_error", True),
    )
    for key, attr in (("error", "error"), ("error_title", "errorTitle"), ("prompt", "prompt"), ("prompt_title", "promptTitle"), ("operator", "operator")):
        if key in spec and spec[key] is not None:
            setattr(dv, attr, spec[key])
    ws.add_data_validation(dv)
    dv.add(ref)


def _add_conditional(ws, ref: str, spec: dict) -> None:
    if not isinstance(spec, dict):
        raise ValueError(f"invalid conditional spec: {spec!r}")
    rule = spec.get("rule", "cell_is")
    dxf = DifferentialStyle(
        fill=PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    )
    if rule == "cell_is":
        ws.conditional_formatting.add(
            ref,
            CellIsRule(
                operator=spec.get("operator", "greaterThan"),
                formula=spec.get("formula", ["0"]),
                fill=dxf.fill,
            ),
        )
    elif rule == "expression":
        if "formula" not in spec:
            raise ValueError("conditional 'expression' rule requires 'formula'")
        ws.conditional_formatting.add(ref, FormulaRule(formula=spec["formula"], fill=dxf.fill))
    else:
        raise ValueError(f"unsupported conditional rule: {rule!r}")


def format_range(path: str, sheet: str, cell_range: str, style: dict) -> dict:
    if not isinstance(style, dict):
        raise ValueError(f"style must be a dict, got {type(style).__name__}")
    unknown = [key for key in style if key not in STYLE_KEYS]
    if unknown:
        raise ValueError(f"invalid style keys: {unknown!r}")

    bounds = _parse_range(cell_range)
    ref = _range_ref(bounds)

    body_font = _parse_font(resolve_style("body_font", None))

    font = None
    if "font" in style and style["font"] is not None:
        font = _to_font(_merge_font(body_font, _parse_font(style["font"])))

    fill = None
    if "fill" in style and style["fill"] is not None:
        fill = PatternFill(
            start_color=_colour_code(style["fill"], "fill colour"),
            end_color=_colour_code(style["fill"], "fill colour"),
            fill_type="solid",
        )

    border = None
    if "border" in style and style["border"] is not None:
        border = _parse_border(style["border"])

    number_format = None
    if "number_format" in style and style["number_format"] is not None:
        if not isinstance(style["number_format"], str) or not style["number_format"]:
            raise ValueError(f"invalid number format: {style['number_format']!r}")
        number_format = style["number_format"]

    alignment = None
    if "alignment" in style and style["alignment"] is not None:
        alignment = _parse_alignment(style["alignment"])

    for key in ("col_widths", "row_height", "freeze", "print_fit", "table_style", "validation", "conditional"):
        if key in style and style[key] is not None:
            if key == "row_height":
                _positive_number(style[key], "row height")
            if key == "freeze" and not isinstance(style[key], (str, bool)):
                raise ValueError(f"invalid freeze spec: {style[key]!r}")

    wb = _open_wb(path)
    try:
        ws = _worksheet(wb, sheet)

        _paint_range(ws, bounds, font=font, fill=fill, border=border, number_format=number_format, alignment=alignment)
        _set_dimensions(ws, bounds, col_widths=style.get("col_widths"), row_height=style.get("row_height"))

        if "freeze" in style and style["freeze"] is not None:
            freeze = style["freeze"]
            if freeze is True:
                ws.freeze_panes = resolve_style("freeze_panes", None)
            elif freeze is False:
                ws.freeze_panes = None
            else:
                ws.freeze_panes = freeze

        if "print_fit" in style and style["print_fit"] is not None:
            _apply_print(ws, style["print_fit"])

        table_name = None
        if "table_style" in style and style["table_style"] is not None:
            table_name = _add_table(ws, ref, style["table_style"])

        if "validation" in style and style["validation"] is not None:
            _add_validation(ws, ref, style["validation"])

        if "conditional" in style and style["conditional"] is not None:
            _add_conditional(ws, ref, style["conditional"])

        wb.save(path)
        result: dict = {"ok": True, "sheet": sheet, "range": ref}
        if table_name is not None:
            result["table"] = table_name
        return result
    finally:
        wb.close()


def apply_ib_table(path: str, sheet: str, cell_range: str, has_totals: bool) -> dict:
    bounds = _parse_range(cell_range)
    min_col, min_row, max_col, max_row = bounds

    header_fill = _colour_code(resolve_style("header_fill", None), "header fill")
    header_font = _to_font(
        _merge_font(
            _parse_font(resolve_style("body_font", None)),
            _parse_font(resolve_style("header_font", None)),
        )
    )
    body_font = _to_font(_parse_font(resolve_style("body_font", None)))
    int_format = resolve_style("number_format_int", None)
    grid = Side(
        style=resolve_style("border_grid_style", None),
        color=_colour_code(resolve_style("border_grid", None), "grid colour"),
    )
    thin_border = Border(left=grid, right=grid, top=grid, bottom=grid)
    total_side = Side(
        style=resolve_style("border_total_style", None),
        color=_colour_code(resolve_style("border_grid", None), "grid colour"),
    )

    wb = _open_wb(path)
    try:
        ws = _worksheet(wb, sheet)

        header_bounds = (min_col, min_row, max_col, min_row)
        _paint_range(
            ws,
            header_bounds,
            font=header_font,
            fill=PatternFill(start_color=header_fill, end_color=header_fill, fill_type="solid"),
            border=thin_border,
        )

        last_body_row = max_row - 1 if has_totals else max_row
        if last_body_row >= min_row + 1:
            _paint_range(
                ws,
                (min_col, min_row + 1, max_col, last_body_row),
                font=body_font,
                border=thin_border,
                number_format=int_format,
            )

        if has_totals:
            total_border = Border(left=grid, right=grid, top=grid, bottom=total_side)
            _paint_range(
                ws,
                (min_col, max_row, max_col, max_row),
                font=body_font,
                border=total_border,
                number_format=int_format,
            )

        width = resolve_style("col_width", None)
        for offset in range(max_col - min_col + 1):
            ws.column_dimensions[get_column_letter(min_col + offset)].width = width

        height = resolve_style("row_height", None)
        for row in range(min_row, max_row + 1):
            ws.row_dimensions[row].height = height

        ws.freeze_panes = resolve_style("freeze_panes", None)

        print_setup = resolve_style("print_setup", None)
        ws.page_setup.orientation = print_setup["orientation"]
        ws.page_setup.fitToWidth = print_setup["fit_to_pages_wide"]
        ws.page_setup.fitToHeight = print_setup["fit_to_pages_tall"]
        footer = str(print_setup["footer"])
        ws.oddFooter.left.text = None
        ws.oddFooter.center.text = footer
        ws.oddFooter.right.text = None
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

        wb.save(path)
        return {"ok": True, "sheet": sheet, "range": _range_ref(bounds), "has_totals": has_totals}
    finally:
        wb.close()
