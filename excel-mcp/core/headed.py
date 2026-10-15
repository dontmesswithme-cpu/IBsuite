import time
from contextlib import contextmanager
from pathlib import Path

from openpyxl.utils import get_column_letter, range_boundaries

CALL_REJECTED = -2147418111
SERVERCALL_RETRYLATER = -2147417846
_RETRY_DELAYS = (0.2, 0.5, 1.0)
MAX_HEADED_CELLS = 100000


class HeadedUnavailable(Exception):
    pass


def _parse_range(value: str, what: str) -> tuple[int, int, int, int]:
    text = (value or "").strip().replace("$", "")
    if not text:
        raise ValueError(f"bad {what}: {value!r}")
    try:
        return range_boundaries(text)
    except Exception as exc:
        raise ValueError(f"bad {what}: {value!r}") from exc


def _range_ref(bounds: tuple[int, int, int, int]) -> str:
    min_col, min_row, max_col, max_row = bounds
    start = f"{get_column_letter(min_col)}{min_row}"
    end = f"{get_column_letter(max_col)}{max_row}"
    return start if start == end else f"{start}:{end}"


def _require_file(path: str) -> Path:
    target = Path(path)
    if not target.is_file():
        raise FileNotFoundError(f"workbook not found: {target}")
    return target


def _check_values(values) -> list[list]:
    if not isinstance(values, (list, tuple)) or len(values) == 0:
        raise ValueError("values must be a non-empty list of rows")
    rows = []
    for row in values:
        if not isinstance(row, (list, tuple)) or len(row) == 0:
            raise ValueError("values must be a non-empty list of rows")
        rows.append(list(row))
    width = len(rows[0])
    for row in rows:
        if len(row) != width:
            raise ValueError("ragged values: every row must have the same length")
    total = len(rows) * width
    if total > MAX_HEADED_CELLS:
        raise ValueError(
            f"write limited to {MAX_HEADED_CELLS} cells per call, split into batches (got {total})"
        )
    return rows


def _launch():
    try:
        import win32com.client
    except ImportError as exc:
        raise HeadedUnavailable("pywin32 not installed") from exc
    try:
        return win32com.client.Dispatch("Excel.Application")
    except Exception as exc:
        raise HeadedUnavailable(f"Excel.Application unavailable: {exc}") from exc


@contextmanager
def _excel():
    app = _launch()
    try:
        app.Visible = True
        app.DisplayAlerts = False
        app.ScreenUpdating = False
        yield app
    finally:
        try:
            app.Quit()
        except Exception:
            pass


def _com_error_type():
    try:
        import pywintypes
        return pywintypes.com_error
    except ImportError:
        return None


def _is_busy_or_locked(exc, com_error) -> bool:
    if com_error is not None and isinstance(exc, com_error):
        hresult = getattr(exc, "hresult", None)
        if hresult in (CALL_REJECTED, SERVERCALL_RETRYLATER):
            return True
    if isinstance(exc, PermissionError):
        return True
    msg = str(exc).lower()
    return "call was rejected" in msg or "busy" in msg or "locked" in msg or "retry" in msg


def _restore_backup_if_any(path: str, ops: dict | None = None) -> None:
    try:
        from core.versions import list_versions, rollback

        b_id = ops.get("backup_id") if ops else None
        if not b_id:
            all_b = list_versions(path)
            if all_b:
                b_id = all_b[-1]
        if b_id:
            rollback(path, b_id)
    except Exception:
        pass


def _with_retry(operation, on_retry_exhausted=None):
    com_error = _com_error_type()
    last_rejected = None
    for index, delay in enumerate((0.0,) + _RETRY_DELAYS):
        if delay:
            time.sleep(delay)
        try:
            return operation()
        except HeadedUnavailable:
            raise
        except Exception as exc:
            if _is_busy_or_locked(exc, com_error):
                last_rejected = exc
                if index < len(_RETRY_DELAYS):
                    continue
                break
            if com_error is not None and isinstance(exc, com_error):
                raise HeadedUnavailable(f"Excel COM error: {exc}") from exc
            raise
    if last_rejected is not None:
        if on_retry_exhausted:
            return on_retry_exhausted(last_rejected)
        return {
            "ok": False,
            "status": "needs_retry",
            "reason": f"file locked or Excel busy after retries: {last_rejected}",
        }
    raise HeadedUnavailable(f"Excel busy/detached after retries: {last_rejected}") from last_rejected


def _com_sheet(wb, sheet: str):
    try:
        names = [ws.Name for ws in wb.Worksheets]
    except Exception as exc:
        raise HeadedUnavailable(f"cannot list sheets: {exc}") from exc
    if sheet not in names:
        raise ValueError(f"unknown sheet: {sheet!r}")
    return wb.Worksheets(sheet)


def _to_grid(value, rows: int, cols: int) -> list[list]:
    if rows == 1 and cols == 1:
        return [[value]]
    if not isinstance(value, (tuple, list)):
        value = ((value,),)
    grid: list[list] = []
    for r in range(rows):
        if r < len(value) and isinstance(value[r], (tuple, list)):
            row = list(value[r])
        elif r < len(value):
            row = [value[r]]
        else:
            row = []
        while len(row) < cols:
            row.append(None)
        grid.append(row[:cols])
    return grid


def is_excel_available() -> bool:
    try:
        with _excel():
            return True
    except HeadedUnavailable:
        return False
    except Exception:
        return False


def headed_read(path: str, sheet: str, cell_range: str) -> dict:
    bounds = _parse_range(cell_range, "cell range")
    _require_file(path)
    min_col, min_row, max_col, max_row = bounds
    rows = max_row - min_row + 1
    cols = max_col - min_col + 1
    ref = _range_ref(bounds)

    def _do():
        with _excel() as app:
            wb = app.Workbooks.Open(str(Path(path).resolve()), ReadOnly=True)
            try:
                ws = _com_sheet(wb, sheet)
                rng = ws.Range(ref)
                values = _to_grid(rng.Value, rows, cols)
                formulas = _to_grid(rng.Formula, rows, cols)
                formats = _to_grid(rng.NumberFormat, rows, cols)
                return {"values": values, "formulas": formulas, "number_formats": formats}
            finally:
                wb.Close(SaveChanges=False)

    return _with_retry(_do)


def headed_apply(path: str, ops: dict) -> dict:
    tool = ops.get("tool", "write_range")
    if tool not in ("write_range", "set_formula", "add_sheet"):
        raise HeadedUnavailable(f"headed_apply does not implement {tool!r}")

    sheet = ops.get("sheet", "")
    if tool == "write_range":
        for key in ("sheet", "cell_range", "values"):
            if key not in ops:
                raise ValueError(f"write_range missing param: {key!r}")
        rows = _check_values(ops["values"])
        bounds = _parse_range(ops["cell_range"], "cell range")
        _require_file(path)
        ref = _range_ref(bounds)
        payload = rows[0][0] if len(rows) == 1 and len(rows[0]) == 1 else tuple(tuple(r) for r in rows)

        def _on_retry_exhausted(exc):
            _restore_backup_if_any(path, ops)
            return {
                "ok": False,
                "status": "needs_retry",
                "reason": f"file locked or Excel busy after retries: {exc}",
            }

        # Save on success, discard on error; always release the workbook.
        def _do_write_clean():
            with _excel() as app:
                wb = app.Workbooks.Open(str(Path(path).resolve()))
                try:
                    ws = _com_sheet(wb, sheet)
                    ws.Range(ref).Value = payload
                    wb.Save()
                    result = {"ok": True, "sheet": sheet, "range": ref, "rows": len(rows), "cols": len(rows[0])}
                except Exception:
                    _restore_backup_if_any(path, ops)
                    raise
                finally:
                    try:
                        wb.Close(SaveChanges=False)
                    except Exception:
                        pass
                return result

        return _with_retry(_do_write_clean, on_retry_exhausted=_on_retry_exhausted)

    if tool == "set_formula":
        for key in ("sheet", "cell", "formula"):
            if key not in ops:
                raise ValueError(f"set_formula missing param: {key!r}")
        formula = ops["formula"]
        if not isinstance(formula, str) or not formula.strip().startswith("="):
            raise ValueError(f"formula must start with '=': {formula!r}")
        bounds = _parse_range(ops["cell"], "cell range")
        if bounds[0] != bounds[2] or bounds[1] != bounds[3]:
            raise ValueError(f"bad cell range: single cell expected, got {ops['cell']!r}")
        _require_file(path)
        ref = _range_ref(bounds)

        def _on_formula_exhausted(exc):
            _restore_backup_if_any(path, ops)
            return {
                "ok": False,
                "status": "needs_retry",
                "reason": f"file locked or Excel busy after retries: {exc}",
            }

        def _do_formula():
            with _excel() as app:
                wb = app.Workbooks.Open(str(Path(path).resolve()))
                try:
                    ws = _com_sheet(wb, sheet)
                    ws.Range(ref).Formula = formula
                    wb.Save()
                    result = {"ok": True, "sheet": sheet, "cell": ref, "formula": formula}
                except Exception:
                    _restore_backup_if_any(path, ops)
                    raise
                finally:
                    try:
                        wb.Close(SaveChanges=False)
                    except Exception:
                        pass
                return result

        return _with_retry(_do_formula, on_retry_exhausted=_on_formula_exhausted)

    for key in ("name",):
        if key not in ops:
            raise ValueError(f"add_sheet missing param: {key!r}")
    title = (ops["name"] or "").strip()
    if not title:
        raise ValueError("empty sheet name")
    _require_file(path)

    def _on_add_exhausted(exc):
        _restore_backup_if_any(path, ops)
        return {
            "ok": False,
            "status": "needs_retry",
            "reason": f"file locked or Excel busy after retries: {exc}",
        }

    def _do_add():
        with _excel() as app:
            wb = app.Workbooks.Open(str(Path(path).resolve()))
            try:
                try:
                    existing = [ws.Name for ws in wb.Worksheets]
                except Exception as exc:
                    raise HeadedUnavailable(f"cannot list sheets: {exc}") from exc
                if title in existing:
                    raise ValueError(f"sheet already exists: {title!r}")
                wb.Worksheets.Add().Name = title
                wb.Save()
                result = {"ok": True, "sheet": title}
            except Exception:
                _restore_backup_if_any(path, ops)
                raise
            finally:
                try:
                    wb.Close(SaveChanges=False)
                except Exception:
                    pass
            return result

    return _with_retry(_do_add, on_retry_exhausted=_on_add_exhausted)
