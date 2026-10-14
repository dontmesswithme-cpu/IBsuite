"""Task 3 tests: copy-before-write backups, FIFO retention, byte-identical rollback, audit log."""

import hashlib
import json

from openpyxl import Workbook, load_workbook

import core.audit as audit
import core.versions as versions
from core.audit import append
from core.versions import backup, list_versions, rollback

MARK = "orig"


def make_minimal_xlsx(path: str) -> None:
    wb = Workbook()
    wb.active["A1"] = MARK
    wb.save(path)


def corrupt(path: str) -> None:
    wb = Workbook()
    wb.active["A1"] = "CLOBBERED"
    wb.save(path)


def read_back(path: str) -> str:
    return load_workbook(path)["Sheet"]["A1"].value


def sha(path: str) -> str:
    with open(path, "rb") as fh:
        return hashlib.sha256(fh.read()).hexdigest()


def test_backup_restore_on_apply_crash(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))
    original_sha = sha(str(f))

    bid = backup(str(f), "write_range", "A1")
    assert bid
    assert len(list_versions(str(f))) == 1

    corrupt(str(f))
    assert read_back(str(f)) != MARK

    result = rollback(str(f), bid)
    assert result["ok"] is True
    assert read_back(str(f)) == MARK
    assert sha(str(f)) == original_sha


def test_fifo_eviction(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    for i in range(21):
        backup(str(f), "t", f"r{i}")

    assert len(list_versions(str(f))) == 20


def test_fifo_evicts_oldest_first(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    ids = [backup(str(f), "t", f"r{i}") for i in range(21)]
    retained = list_versions(str(f))

    assert ids[0] not in retained
    assert set(retained) == set(ids[1:])


def test_rollback_backs_up_pre_rollback_state(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    bid = backup(str(f), "write_range", "A1")
    corrupt(str(f))

    rollback(str(f), bid)

    versions_for_file = list_versions(str(f))
    assert bid in versions_for_file
    assert len(versions_for_file) == 2


def test_backup_writes_json_sidecar(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    bid = backup(str(f), "write_range", "A1:A2")
    sidecar = tmp_path / "m.versions" / (bid[:-4] + ".json")

    assert sidecar.exists()
    meta = json.loads(sidecar.read_text(encoding="utf-8"))
    assert meta["tool"] == "write_range"
    assert meta["label"] == "A1:A2"
    assert meta["backup_id"] == bid


def test_rollback_unknown_backup_id_rejected(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    result = rollback(str(f), "not-a-real-backup.bak")
    assert result["ok"] is False


def test_versions_dir_is_stem_sibling(tmp_path):
    f = tmp_path / "model.xlsx"
    make_minimal_xlsx(str(f))

    backup(str(f), "write_range", "A1")

    assert (tmp_path / "model.versions").is_dir()
    assert (tmp_path / "model.versions").parent == tmp_path


def test_audit_append_writes_pipe_delimited_line(tmp_path, monkeypatch):
    log = tmp_path / "audit.log"
    monkeypatch.setattr(audit, "AUDIT_LOG_PATH", log)

    append("write_range", "model.xlsx", "Sheet1:A1", "bid123", "ok")

    fields = [f.strip() for f in log.read_text(encoding="utf-8").strip().split("|")]
    assert fields[1] == "write_range"
    assert fields[2] == "model.xlsx"
    assert fields[3] == "Sheet1:A1"
    assert fields[4] == "bid123"
    assert fields[5] == "ok"


def test_audit_append_creates_log_and_appends(tmp_path, monkeypatch):
    log = tmp_path / "audit.log"
    monkeypatch.setattr(audit, "AUDIT_LOG_PATH", log)

    append("read_range", "model.xlsx", "Sheet1:A1", None, "ok")
    append("write_range", "model.xlsx", "Sheet1:A2", "b2", "ERROR")

    lines = log.read_text(encoding="utf-8").strip().splitlines()
    assert len(lines) == 2
    assert lines[0].split("|")[4].strip() == "-"
    assert "ERROR" in lines[1]


def test_audit_passes_through_caller_redacted_detail(tmp_path, monkeypatch):
    log = tmp_path / "audit.log"
    monkeypatch.setattr(audit, "AUDIT_LOG_PATH", log)

    macro = "Shell('secret-token')"
    redacted = "sha256:" + hashlib.sha256(macro.encode()).hexdigest()[:16]
    append("macro_run_gated", "m.xlsm", redacted, "b3", "ok")

    written = log.read_text(encoding="utf-8")
    assert redacted in written
    assert "secret-token" not in written


def test_audit_line_is_append_only(tmp_path, monkeypatch):
    log = tmp_path / "audit.log"
    monkeypatch.setattr(audit, "AUDIT_LOG_PATH", log)

    for i in range(5):
        append("read_range", "model.xlsx", f"Sheet1:A{i}", None, "ok")

    assert len(log.read_text(encoding="utf-8").strip().splitlines()) == 5


def test_audit_default_path_is_repo_audit_log():
    assert audit.AUDIT_LOG_PATH.name == "audit.log"
    assert "excel-mcp" in str(audit.AUDIT_LOG_PATH)


def test_versions_retention_constant_is_twenty():
    assert versions.MAX_VERSIONS == 20
