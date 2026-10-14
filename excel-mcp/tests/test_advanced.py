"""Task 6 tests: headless charts, pivots (SUMIFS fallback), images."""

import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image as PILImage

from core.advanced_headless import add_chart, add_image, add_pivot


def make_table_xlsx(path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Month"
    ws["B1"] = "Sales"
    for i, (month, sales) in enumerate(
        [("Jan", 100), ("Feb", 140), ("Mar", 90), ("Apr", 160)], start=2
    ):
        ws.cell(row=i, column=1).value = month
        ws.cell(row=i, column=2).value = sales
    wb.save(path)


def make_png(path: str) -> None:
    img = PILImage.new("RGB", (120, 60), color=(0, 32, 96))
    img.save(path)


def has_drawing(path: str, sheet: str) -> bool:
    wb = load_workbook(path)
    try:
        ws = wb[sheet]
        return bool(ws._charts or ws._images)
    finally:
        wb.close()


# --- plan Step 1 contract test (verbatim scope) ---


def test_chart_and_image_roundtrip(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))
    add_chart(str(f), "Sheet1", "A1:B5", "line", "E1")
    img = tmp_path / "p.png"
    make_png(str(img))
    add_image(str(f), "Sheet1", str(img), "E15")
    assert has_drawing(str(f), "Sheet1") is True


# --- hardening ---


def test_chart_types_and_return_shape(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))

    for chart_type in ("line", "bar", "column", "pie"):
        out = add_chart(str(f), "Sheet1", "A1:B5", chart_type, "E1")
        assert out["ok"] is True and out["chart"] == chart_type


def test_chart_rejects_unknown_type_and_sheet(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))
    before = f.read_bytes()

    with pytest.raises(ValueError, match="unknown chart type"):
        add_chart(str(f), "Sheet1", "A1:B5", "hologram", "E1")
    with pytest.raises(ValueError, match="unknown sheet"):
        add_chart(str(f), "Ghost", "A1:B5", "line", "E1")
    with pytest.raises(ValueError, match="bad data range"):
        add_chart(str(f), "Sheet1", "junk", "line", "E1")

    assert f.read_bytes() == before


def test_add_image_rejects_missing_file_and_bad_anchor(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))
    before = f.read_bytes()

    with pytest.raises(FileNotFoundError):
        add_image(str(f), "Sheet1", str(tmp_path / "nope.png"), "E15")
    with pytest.raises(ValueError, match="anchor"):
        add_image(str(f), "Sheet1", str(tmp_path / "p.png"), "junk!!")

    assert f.read_bytes() == before


def test_add_image_returns_ok_with_anchor(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))
    img = tmp_path / "p.png"
    make_png(str(img))

    out = add_image(str(f), "Sheet1", str(img), "E15")
    assert out["ok"] is True and out["anchor"] == "E15"


def test_pivot_fallback_builds_sumifs_summary(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))

    out = add_pivot(
        str(f), "Sheet1", "A1:B5", "Pivot1", rows=["Month"], values=["Sales"]
    )
    assert out["ok"] is True
    assert "warning" in out

    wb = load_workbook(str(f))
    try:
        assert "Pivot1" in wb.sheetnames
        ws = wb["Pivot1"]
        assert ws["A1"].value == "Month"
        assert ws["B1"].value == "SUM of Sales"
        assert "SUMIFS" in str(ws["B2"].value)
    finally:
        wb.close()


def test_pivot_rejects_bad_source_and_unknown_sheet(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))
    before = f.read_bytes()

    with pytest.raises(ValueError, match="unknown sheet"):
        add_pivot(str(f), "Ghost", "A1:B5", "P1", rows=["Month"], values=["Sales"])
    with pytest.raises(ValueError, match="source range"):
        add_pivot(str(f), "Sheet1", "junk", "P1", rows=["Month"], values=["Sales"])

    assert f.read_bytes() == before


def test_pivot_rejects_existing_dest_sheet(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))

    with pytest.raises(ValueError, match="already exists"):
        add_pivot(str(f), "Sheet1", "A1:B5", "Sheet1", rows=["Month"], values=["Sales"])
