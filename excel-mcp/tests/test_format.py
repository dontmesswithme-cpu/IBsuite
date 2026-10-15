"""Task 5 tests: IB-grade headless formatting against the golden file."""

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from core.format_headless import apply_ib_table, format_range

GOLDEN = json.loads((Path(__file__).parent / "golden_ib.json").read_text(encoding="utf-8"))


def make_minimal_xlsx(path: str) -> None:
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = "Revenue"
    wb.active["B1"] = "2024"
    wb.active["C1"] = "2025"
    wb.active["A2"] = "Sales"
    wb.active["B2"] = 100
    wb.active["C2"] = 120
    wb.active["A3"] = "Total"
    wb.active["B3"] = 100
    wb.active["C3"] = 120
    wb.save(path)


def _rgb(color) -> str | None:
    if color is None or color.rgb is None:
        return None
    rgb = str(color.rgb)
    return rgb[2:] if len(rgb) == 8 else rgb


def inspect_fills(path: str, sheet: str, cell_range: str) -> list[list[str | None]]:
    from openpyxl.utils import range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    wb = load_workbook(path)
    try:
        ws = wb[sheet]
        return [
            [_rgb(ws.cell(row=r, column=c).fill.start_color) for c in range(min_col, max_col + 1)]
            for r in range(min_row, max_row + 1)
        ]
    finally:
        wb.close()


def inspect_formats(path: str, sheet: str, cell_range: str) -> list[list[str]]:
    from openpyxl.utils import range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    wb = load_workbook(path)
    try:
        ws = wb[sheet]
        return [
            [ws.cell(row=r, column=c).number_format for c in range(min_col, max_col + 1)]
            for r in range(min_row, max_row + 1)
        ]
    finally:
        wb.close()


def inspect_sheet(path: str, sheet: str):
    wb = load_workbook(path)
    try:
        return wb[sheet]
    finally:
        wb.close()


# --- plan Step 1 contract test (verbatim scope) ---


def test_ib_header_and_accounting(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))
    apply_ib_table(str(f), "Sheet1", "A1:C3", True)
    fills = inspect_fills(str(f), "Sheet1", "A1:C1")
    assert all(v == "002060" for v in fills[0])
    fmts = inspect_formats(str(f), "Sheet1", "A2:C3")
    assert " #,##0" in fmts[0][0]


# --- golden assertions: every IB marker pinned to golden_ib.json ---


def test_golden_schema_frozen():
    assert set(GOLDEN.keys()) == {"fills", "number_formats", "fonts", "borders", "layout"}
    import re
    src = Path(__file__).read_text(encoding="utf-8")
    accessed_keys = set(re.findall(r'GOLDEN\["([^"]+)"\]', src))
    assert accessed_keys == {"fills", "number_formats", "fonts", "borders", "layout"}


def test_golden_header_fill_and_font(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))
    apply_ib_table(str(f), "Sheet1", "A1:C3", True)

    fills = inspect_fills(str(f), "Sheet1", "A1:C1")
    assert fills[0] == [GOLDEN["fills"]["header"]] * 3

    wb = load_workbook(str(f))
    try:
        font = wb["Sheet1"]["A1"].font
        assert font.bold is GOLDEN["fonts"]["header"]["bold"]
        assert font.size == GOLDEN["fonts"]["header"]["size"]
        assert _rgb(font.color) == GOLDEN["fonts"]["header"]["color"]
        assert font.name == GOLDEN["fonts"]["header"]["name"]
    finally:
        wb.close()


def test_golden_body_number_format_and_font(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))
    apply_ib_table(str(f), "Sheet1", "A1:C3", True)

    fmts = inspect_formats(str(f), "Sheet1", "A2:C3")
    for row in fmts:
        assert row == [GOLDEN["number_formats"]["body"]] * 3

    wb = load_workbook(str(f))
    try:
        font = wb["Sheet1"]["B2"].font
        assert font.name == GOLDEN["fonts"]["body"]["name"]
        assert font.size == GOLDEN["fonts"]["body"]["size"]
        assert font.bold is GOLDEN["fonts"]["body"]["bold"]
    finally:
        wb.close()


def test_golden_grid_and_total_borders(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))
    apply_ib_table(str(f), "Sheet1", "A1:C3", True)

    wb = load_workbook(str(f))
    try:
        ws = wb["Sheet1"]
        body = ws["B2"].border
        assert body.bottom.style == GOLDEN["borders"]["grid"]["style"]
        assert _rgb(body.bottom.color) == GOLDEN["borders"]["grid"]["color"]

        total = ws["B3"].border
        assert total.bottom.style == GOLDEN["borders"]["total"]["style"]
        assert _rgb(total.bottom.color) == GOLDEN["borders"]["total"]["color"]
    finally:
        wb.close()


def test_golden_freeze_widths_height_and_print(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))
    apply_ib_table(str(f), "Sheet1", "A1:C3", True)

    wb = load_workbook(str(f))
    try:
        ws = wb["Sheet1"]
        assert ws.freeze_panes == GOLDEN["layout"]["freeze_panes"]
        for col in ("A", "B", "C"):
            assert ws.column_dimensions[col].width == GOLDEN["layout"]["col_width"]
        for row in (1, 2, 3):
            assert ws.row_dimensions[row].height == GOLDEN["layout"]["row_height"]
        assert ws.page_setup.orientation == GOLDEN["layout"]["print"]["orientation"]
        assert ws.page_setup.fitToWidth == GOLDEN["layout"]["print"]["fit_to_pages_wide"]
        assert ws.page_setup.fitToHeight == GOLDEN["layout"]["print"]["fit_to_pages_tall"]
        assert ws.oddFooter.center.text == GOLDEN["layout"]["print"]["footer"]
        assert ws.oddFooter.left.text == ""
        assert ws.oddFooter.right.text == ""
        assert ws.oddFooter.center.text == "&F | &A"
        assert f"{ws.oddFooter.left.text}{ws.oddFooter.center.text}{ws.oddFooter.right.text}" == "&F | &A"
    finally:
        wb.close()


def test_apply_ib_table_returns_ok(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    out = apply_ib_table(str(f), "Sheet1", "A1:C3", True)
    assert out["ok"] is True


def test_apply_ib_table_without_totals_has_no_double_border(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))
    apply_ib_table(str(f), "Sheet1", "A1:C3", False)

    wb = load_workbook(str(f))
    try:
        assert wb["Sheet1"]["B3"].border.bottom.style == GOLDEN["borders"]["grid"]["style"]
    finally:
        wb.close()


def test_format_range_applies_explicit_style(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    out = format_range(
        str(f),
        "Sheet1",
        "A1:B1",
        {
            "fill": "FF0000",
            "font": {"bold": True, "size": 12},
            "number_format": "0.00",
            "alignment": {"horizontal": "center"},
        },
    )
    assert out["ok"] is True

    wb = load_workbook(str(f))
    try:
        cell = wb["Sheet1"]["A1"]
        assert _rgb(cell.fill.start_color) == "FF0000"
        assert cell.font.bold is True and cell.font.size == 12
        assert cell.number_format == "0.00"
        assert cell.alignment.horizontal == "center"
        assert wb["Sheet1"]["C1"].number_format == "General"
    finally:
        wb.close()


def test_format_range_unknown_key_rejected(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    with pytest.raises(ValueError, match="style keys"):
        format_range(str(f), "Sheet1", "A1", {"neon_glow": True})


def test_format_range_rejects_bad_range_and_unknown_sheet(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))
    before = f.read_bytes()

    with pytest.raises(ValueError, match="bad cell range"):
        format_range(str(f), "Sheet1", "junk", {"fill": "FF0000"})
    with pytest.raises(ValueError, match="unknown sheet"):
        format_range(str(f), "Ghost", "A1", {"fill": "FF0000"})

    assert f.read_bytes() == before


def test_format_range_freeze_and_widths(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    format_range(
        str(f), "Sheet1", "A1:C3", {"freeze": "C3", "col_widths": {"B": 22}, "row_height": 20}
    )

    wb = load_workbook(str(f))
    try:
        ws = wb["Sheet1"]
        assert ws.freeze_panes == "C3"
        assert ws.column_dimensions["B"].width == 22
        assert ws.row_dimensions[2].height == 20
    finally:
        wb.close()


def test_format_range_table_style_creates_table(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    format_range(str(f), "Sheet1", "A1:C3", {"table_style": "TableStyleMedium9"})

    wb = load_workbook(str(f))
    try:
        ws = wb["Sheet1"]
        assert len(ws.tables) == 1
        table = next(iter(ws.tables.values()))
        assert table.ref == "A1:C3"
        assert table.tableStyleInfo.name == "TableStyleMedium9"
    finally:
        wb.close()


def test_format_range_validation_and_conditional(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    format_range(
        str(f),
        "Sheet1",
        "B2:B3",
        {
            "validation": {"type": "whole", "operator": "between", "formula1": "0", "formula2": "1000"},
            "conditional": {"operator": "greaterThan", "formula": ["500"], "fill": "FFC7CE"},
        },
    )

    wb = load_workbook(str(f))
    try:
        ws = wb["Sheet1"]
        assert len(ws.data_validations.dataValidation) == 1
        assert len(ws.conditional_formatting._cf_rules) == 1
    finally:
        wb.close()
