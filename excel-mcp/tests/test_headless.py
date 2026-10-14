"""Task 4 tests: headless list/read/write/formula/sheets and read-only preview_diff."""

import pytest
from openpyxl import Workbook

from core.headless import (
    MAX_CELLS,
    add_sheet,
    list_sheets,
    preview_diff,
    read_range,
    set_formula,
    write_range,
)


def make_minimal_xlsx(path: str) -> None:
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.save(path)


def make_seeded_xlsx(path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Revenue"
    ws["B1"] = 100
    ws["C1"] = 200
    ws["A2"] = "COGS"
    ws["B2"] = 40
    ws["C2"] = 60
    wb.create_sheet("Notes")
    wb.save(path)


# --- plan Step 1 contract tests (verbatim scope) ---


def test_oversize_rejected_untouched(tmp_path):
    f = tmp_path / "big.xlsx"
    make_minimal_xlsx(str(f))
    before = f.read_bytes()
    with pytest.raises(ValueError, match="split"):
        write_range(str(f), "Sheet1", "A1", [[1]] * 200000)
    assert f.read_bytes() == before


def test_roundtrip_values_formulas(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))
    write_range(str(f), "Sheet1", "A1:B2", [[1, 2], [3, 4]])
    set_formula(str(f), "Sheet1", "C1", "=SUM(A1:B1)")
    out = read_range(str(f), "Sheet1", "A1:C1")
    assert out["values"][0][:2] == [1, 2]
    assert out["formulas"][0][2] == "=SUM(A1:B1)"


def test_preview_diff_never_writes(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))
    before = f.read_bytes()
    d = preview_diff(
        str(f),
        {"tool": "write_range", "sheet": "Sheet1", "cell_range": "A1", "values": [[99]]},
    )
    assert "A1" in str(d["diff"]) and f.read_bytes() == before


# --- hardening beyond the plan's three ---


def test_max_cells_constant_is_100k():
    assert MAX_CELLS == 100000


def test_write_range_returns_rows_and_cols(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    out = write_range(str(f), "Sheet1", "A1", [[10, 20, 30], [40, 50, 60]])
    assert out["ok"] is True and out["rows"] == 2 and out["cols"] == 3

    reread = read_range(str(f), "Sheet1", "A1:C2")
    assert reread["values"] == [[10, 20, 30], [40, 50, 60]]


def test_shape_mismatch_rejected_untouched(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))
    before = f.read_bytes()

    with pytest.raises(ValueError, match="expects"):
        write_range(str(f), "Sheet1", "A1:C3", [[1, 2]])

    assert f.read_bytes() == before


def test_unknown_sheet_rejected_untouched(tmp_path):
    f = tmp_path / "m.xlsx"
    make_seeded_xlsx(str(f))
    before = f.read_bytes()

    with pytest.raises(ValueError, match="unknown sheet"):
        write_range(str(f), "Ghost", "A1", [[1]])

    assert f.read_bytes() == before


def test_bad_range_rejected(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    with pytest.raises(ValueError, match="bad cell range"):
        read_range(str(f), "Sheet1", "not-a-range")

    with pytest.raises(ValueError, match="bad cell range"):
        read_range(str(f), "Sheet1", "")


def test_missing_file_is_not_created(tmp_path):
    f = tmp_path / "ghost.xlsx"

    with pytest.raises(FileNotFoundError):
        write_range(str(f), "Sheet1", "A1", [[1]])

    assert not f.exists()


def test_list_sheets_reports_name_dims_frozen_print(tmp_path):
    f = tmp_path / "m.xlsx"
    make_seeded_xlsx(str(f))

    sheets = list_sheets(str(f))
    assert [s["name"] for s in sheets] == ["Sheet1", "Notes"]

    first = sheets[0]
    assert first["dims"] == "A1:C2"
    assert "frozen" in first
    assert "print" in first
    assert "orientation" in first["print"]


def test_add_sheet_creates_then_rejects_dup_and_blank(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    out = add_sheet(str(f), "Model")
    assert out["ok"] is True
    assert "Model" in [s["name"] for s in list_sheets(str(f))]

    with pytest.raises(ValueError, match="already exists"):
        add_sheet(str(f), "Model")

    with pytest.raises(ValueError, match="empty sheet name"):
        add_sheet(str(f), "   ")


def test_set_formula_requires_leading_equals(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    with pytest.raises(ValueError, match="must start with"):
        set_formula(str(f), "Sheet1", "A1", "SUM(A1:B2)")


def test_set_formula_persists(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    out = set_formula(str(f), "Sheet1", "D4", "=B2*C2")
    assert out["ok"] is True and out["cell"] == "D4"

    assert read_range(str(f), "Sheet1", "D4")["values"][0][0] == "=B2*C2"


def test_preview_diff_reports_before_and_after(tmp_path):
    f = tmp_path / "m.xlsx"
    make_seeded_xlsx(str(f))

    d = preview_diff(
        str(f),
        {"tool": "write_range", "sheet": "Sheet1", "cell_range": "B1", "values": [[999]]},
    )
    assert d["before"]["values"] == [[100]]
    assert d["after"]["values"] == [[999]]
    assert d["diff"][0]["cell"] == "B1"
    assert d["diff"][0]["before"] == 100
    assert d["diff"][0]["after"] == 999


def test_preview_diff_unknown_sheet_does_not_touch_disk(tmp_path):
    f = tmp_path / "m.xlsx"
    make_seeded_xlsx(str(f))
    before = f.read_bytes()

    with pytest.raises(ValueError, match="unknown sheet"):
        preview_diff(
            str(f),
            {"tool": "write_range", "sheet": "Ghost", "cell_range": "A1", "values": [[1]]},
        )

    assert f.read_bytes() == before


def test_preview_diff_is_pure_read_on_oversize_ops(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))
    before = f.read_bytes()

    with pytest.raises(ValueError, match="split"):
        preview_diff(
            str(f),
            {"tool": "write_range", "sheet": "Sheet1", "cell_range": "A1", "values": [[1]] * 200000},
        )

    assert f.read_bytes() == before


def test_read_range_empty_region_returns_empty_grid(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    out = read_range(str(f), "Sheet1", "Z10")
    assert out["values"] == [[None]]
    assert out["formulas"] == [[None]]


def test_read_range_number_formats_present(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))
    write_range(str(f), "Sheet1", "A1", [[1234.5]])

    out = read_range(str(f), "Sheet1", "A1")
    assert out["number_formats"][0][0] == "General"


def test_write_range_accepts_dollar_signs_in_range(tmp_path):
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    write_range(str(f), "Sheet1", "$A$1:$B$1", [[7, 8]])
    assert read_range(str(f), "Sheet1", "A1:B1")["values"][0] == [7, 8]


def test_writes_release_the_file_handle(tmp_path):
    f = tmp_path / "locked.xlsx"
    make_minimal_xlsx(str(f))
    write_range(str(f), "Sheet1", "A1", [[1]])

    with open(f, "r+b"):
        pass
