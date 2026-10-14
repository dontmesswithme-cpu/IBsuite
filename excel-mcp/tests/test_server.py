"""Task 9 tests: paranoid MCP flow (whitelist -> approval -> backup -> dispatch -> audit)."""

import pytest
from openpyxl import Workbook, load_workbook

import config
import core.audit as audit_mod
import server
from server import call_tool, executed_flag, ops_hash


def make_minimal_xlsx(path: str) -> None:
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = 1
    wb.save(path)
    wb.close()


@pytest.fixture
def sandbox(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "ALLOWED_ROOTS", [str(tmp_path)])
    log = tmp_path / "audit.log"
    monkeypatch.setattr(audit_mod, "AUDIT_LOG_PATH", log)
    return tmp_path, log


def read_cell(path: str):
    wb = load_workbook(path)
    try:
        return wb["Sheet1"]["A1"].value
    finally:
        wb.close()


# --- plan Step 1 contract tests (verbatim scope) ---


def test_macro_without_flag_never_executes(tmp_path, monkeypatch):
    log = tmp_path / "audit.log"
    monkeypatch.setattr(audit_mod, "AUDIT_LOG_PATH", log)
    monkeypatch.setattr(config, "ALLOWED_ROOTS", [str(tmp_path)])

    f = tmp_path / "m.xlsm"
    make_minimal_xlsx(str(f))
    before = f.read_bytes()

    out = call_tool("macro_run_gated", {"path": str(f), "macro": "Shell('evil')"})
    assert out["status"] == "needs_approval"
    assert executed_flag() is False
    assert f.read_bytes() == before


def test_write_needs_approval_then_succeeds(sandbox):
    tmp_path, _log = sandbox
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))
    ops = {
        "path": str(f),
        "sheet": "Sheet1",
        "cell_range": "A1",
        "values": [[9]],
        "mode": "headless",
    }
    first = call_tool("write_range", ops)
    assert first["status"] == "needs_approval"
    ops["approval_token"] = first["ops_hash"]
    second = call_tool("write_range", ops)
    assert second["ok"] is True and second["backup_id"]
    assert read_cell(str(f)) == 9


# --- reads need no approval ---


def test_read_tools_need_no_approval(sandbox):
    tmp_path, _log = sandbox
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    assert call_tool("list_sheets", {"path": str(f)})["ok"] is True
    out = call_tool("read_range", {"path": str(f), "sheet": "Sheet1", "cell_range": "A1"})
    assert out["ok"] is True and out["values"] == [[1]]

    before = f.read_bytes()
    diff = call_tool(
        "preview_diff",
        {"path": str(f), "ops": {"tool": "write_range", "sheet": "Sheet1", "cell_range": "A1", "values": [[99]]}},
    )
    assert diff["ok"] is True and "A1" in str(diff["diff"])
    assert f.read_bytes() == before


def test_screenshot_sheet_thumbnail_always_full_on_demand(sandbox):
    tmp_path, _log = sandbox
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    out = call_tool("screenshot_sheet", {"path": str(f), "sheet": "Sheet1", "cell_range": "A1"})
    assert out["ok"] is True and out["thumbnail_b64"] and out["full_b64"] is None

    out2 = call_tool(
        "screenshot_sheet",
        {"path": str(f), "sheet": "Sheet1", "cell_range": "A1", "full": True},
    )
    assert out2["full_b64"] and len(out2["full_b64"]) > len(out2["thumbnail_b64"])


# --- whitelist is first, before backup and audit ---


def test_outside_path_rejected_before_approval_backup_and_audit(sandbox):
    tmp_path, log = sandbox
    outside = "C:\\Windows\\foo.xlsx"

    denied = call_tool("read_range", {"path": outside, "sheet": "S", "cell_range": "A1"})
    assert denied["status"] == "rejected" and denied["ok"] is False

    approved_ops = {
        "path": outside,
        "sheet": "S",
        "cell_range": "A1",
        "values": [[1]],
        "approval_token": ops_hash(
            "write_range",
            {"path": outside, "sheet": "S", "cell_range": "A1", "values": [[1]]},
        ),
    }
    denied_write = call_tool("write_range", approved_ops)
    assert denied_write["status"] == "rejected"

    assert not log.exists(), "rejected paths must not audit-write"


def test_stale_token_rejected_without_touching_disk(sandbox):
    tmp_path, _log = sandbox
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))
    before = f.read_bytes()

    stale = dict(path=str(f), sheet="Sheet1", cell_range="A1", values=[[2]])
    forged = ops_hash("write_range", dict(stale, values=[[999]]))
    out = call_tool("write_range", {**stale, "approval_token": forged})
    assert out["status"] == "needs_approval" and out["ok"] is False
    assert f.read_bytes() == before


def test_ops_hash_ignores_token_and_mode():
    base = {"path": "a.xlsx", "sheet": "S", "cell_range": "A1", "values": [[1]]}
    assert ops_hash("write_range", base) == ops_hash(
        "write_range", {**base, "approval_token": "x", "mode": "headed"}
    )


# --- gated writes audit ---


def test_gated_write_audits_ok_with_backup_id(sandbox):
    tmp_path, log = sandbox
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    ops = {"path": str(f), "sheet": "Sheet1", "cell_range": "A1", "values": [[5]]}
    ops["approval_token"] = ops_hash("write_range", ops)
    out = call_tool("write_range", ops)
    assert out["ok"] is True

    lines = log.read_text(encoding="utf-8").strip().splitlines()
    assert len(lines) == 1
    fields = [field.strip() for field in lines[0].split("|")]
    assert fields[1] == "write_range" and fields[4] == out["backup_id"] and fields[5] == "ok"
    assert out["thumbnail_b64"]


def test_rollback_needs_approval_then_restores(sandbox):
    tmp_path, _log = sandbox
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    write = {"path": str(f), "sheet": "Sheet1", "cell_range": "A1", "values": [[2]]}
    write["approval_token"] = ops_hash("write_range", write)
    first = call_tool("write_range", write)
    assert first["ok"] is True and read_cell(str(f)) == 2

    roll = {"path": str(f), "backup_id": first["backup_id"]}
    assert call_tool("rollback", roll)["status"] == "needs_approval"
    roll["approval_token"] = ops_hash("rollback", roll)
    done = call_tool("rollback", roll)
    assert done["ok"] is True and done["backup_id"]
    assert read_cell(str(f)) == 1


# --- macros ---


def test_macro_approved_but_headless_rejected_without_execution(sandbox):
    tmp_path, _log = sandbox
    f = tmp_path / "m.xlsm"
    make_minimal_xlsx(str(f))

    ops = {"path": str(f), "macro": "MyMacro", "macro_approved": True, "mode": "headless"}
    ops["approval_token"] = ops_hash("macro_run_gated", ops)
    out = call_tool("macro_run_gated", ops)
    assert out["status"] == "rejected" and out["ok"] is False
    assert executed_flag() is False


def test_macro_denial_logs_sha_never_verbatim(sandbox):
    tmp_path, log = sandbox
    f = tmp_path / "m.xlsm"
    make_minimal_xlsx(str(f))

    call_tool("macro_run_gated", {"path": str(f), "macro": "Shell('s3cret-token')"})
    written = log.read_text(encoding="utf-8")
    assert "s3cret-token" not in written
    assert "Shell" not in written


def test_unknown_tool_rejected(sandbox):
    tmp_path, _log = sandbox
    f = tmp_path / "m.xlsx"
    make_minimal_xlsx(str(f))

    out = call_tool("teleport", {"path": str(f)})
    assert out["status"] == "rejected" and out["ok"] is False
