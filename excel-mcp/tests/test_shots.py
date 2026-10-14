"""Task 8 tests: thumbnails and full screenshots as base64 PNG."""

import base64
import io

import pytest
from openpyxl import Workbook
from PIL import Image as PILImage

from core.shots import full_shot, thumbnail


def make_table_xlsx(path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Item"
    ws["B1"] = "Q1"
    ws["C1"] = "Q2"
    ws["A2"] = "Revenue"
    ws["B2"] = 100
    ws["C2"] = 120
    ws["A3"] = "Total"
    ws["B3"] = 300
    ws["C3"] = 360
    ws["A4"] = "Notes"
    ws["B4"] = "x"
    ws["C4"] = "y"
    ws["A5"] = "End"
    ws["B5"] = 1
    ws["C5"] = 2
    wb.save(path)
    wb.close()


def decode(b64: str) -> PILImage.Image:
    return PILImage.open(io.BytesIO(base64.b64decode(b64)))


def test_thumbnail_smaller_than_full(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))
    t = thumbnail(str(f), "Sheet1", "A1:C5")
    full = full_shot(str(f), "Sheet1", "A1:C5", "headless")
    assert len(t) < len(full) and len(t) < 40000


def test_thumbnail_is_320px_valid_png(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))

    img = decode(thumbnail(str(f), "Sheet1", "A1:C5"))
    assert img.format == "PNG"
    assert img.size[0] == 320
    assert img.size[1] > 0


def test_full_shot_is_1600px_valid_png(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))

    img = decode(full_shot(str(f), "Sheet1", "A1:C5", "headless"))
    assert img.format == "PNG"
    assert img.size[0] == 1600


def test_thumbnail_none_range_uses_used_area(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))

    raw = thumbnail(str(f), "Sheet1", None)
    assert isinstance(raw, str) and len(raw) < 40000
    assert decode(raw).size[0] == 320


def test_thumbnail_reflects_navy_header(tmp_path):
    from core.format_headless import apply_ib_table

    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))
    apply_ib_table(str(f), "Sheet1", "A1:C3", True)

    pixels = decode(thumbnail(str(f), "Sheet1", "A1:C3")).convert("RGB")
    red, green, blue = pixels.getpixel((60, 8))
    assert red < 80 and green < 100 and blue > 40


def test_unknown_sheet_rejected(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))

    with pytest.raises(ValueError, match="unknown sheet"):
        thumbnail(str(f), "Ghost", "A1")


def test_bad_range_rejected(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))

    with pytest.raises(ValueError, match="bad .*range"):
        thumbnail(str(f), "Sheet1", "nope")


def test_missing_file_not_created(tmp_path):
    f = tmp_path / "ghost.xlsx"

    with pytest.raises(FileNotFoundError):
        thumbnail(str(f), "Sheet1", "A1")

    assert not f.exists()


def test_full_shot_unknown_mode_rejected(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))

    with pytest.raises(ValueError, match="mode"):
        full_shot(str(f), "Sheet1", "A1", "hyperdrive")


def test_full_shot_headed_returns_valid_png(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))

    img = decode(full_shot(str(f), "Sheet1", "A1:C5", "headed"))
    assert img.format == "PNG" and img.size[0] > 0


def test_shots_release_the_file_handle(tmp_path):
    f = tmp_path / "m.xlsx"
    make_table_xlsx(str(f))
    thumbnail(str(f), "Sheet1", "A1:C5")
    full_shot(str(f), "Sheet1", "A1:C5", "headless")

    with open(f, "r+b"):
        pass
