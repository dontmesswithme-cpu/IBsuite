import pytest
from docx import Document
import core.audit as audit_mod
from server import call_tool, ops_hash


def make_minimal_docx(path: str) -> None:
    doc = Document()
    doc.add_paragraph("Server Test")
    doc.save(path)


@pytest.fixture
def sandbox(tmp_path, monkeypatch):
    log = tmp_path / "audit.log"
    monkeypatch.setattr(audit_mod, "AUDIT_LOG_PATH", log)
    return tmp_path, log


def test_read_tools_need_no_approval(sandbox):
    tmp_path, _log = sandbox
    f = tmp_path / "d.docx"
    make_minimal_docx(str(f))
    assert call_tool("list_structure", {"path": str(f)})["ok"] is True
    out = call_tool("read_text", {"path": str(f), "start": 0, "end": 0})
    assert out["ok"] is True
    assert out["paragraphs"][0]["text"] == "Server Test"


def test_write_needs_approval_then_succeeds(sandbox):
    tmp_path, _log = sandbox
    f = tmp_path / "d.docx"
    make_minimal_docx(str(f))
    ops = {"path": str(f), "text": "Approved", "style": "Normal"}
    first = call_tool("add_paragraph", ops)
    assert first["status"] == "needs_approval"
    ops["approval_token"] = first["ops_hash"]
    second = call_tool("add_paragraph", ops)
    assert second["ok"] is True and second["backup_id"]


def test_outside_path_rejected_before_approval(sandbox):
    tmp_path, log = sandbox
    outside = "C:\\Windows\\secret.docx"
    denied = call_tool("read_text", {"path": outside, "start": 0, "end": 0})
    assert denied["status"] == "rejected" and denied["ok"] is False
    assert not log.exists()


def test_rollback_needs_approval_then_restores(sandbox):
    tmp_path, _log = sandbox
    f = tmp_path / "d.docx"
    make_minimal_docx(str(f))
    write = {"path": str(f), "text": "New Text", "style": "Normal"}
    write["approval_token"] = ops_hash("add_paragraph", write)
    first = call_tool("add_paragraph", write)
    assert first["ok"] is True
    roll = {"path": str(f), "backup_id": first["backup_id"]}
    assert call_tool("rollback", roll)["status"] == "needs_approval"
    roll["approval_token"] = ops_hash("rollback", roll)
    done = call_tool("rollback", roll)
    assert done["ok"] is True
