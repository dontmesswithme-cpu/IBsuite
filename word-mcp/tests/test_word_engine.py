import pytest
from docx import Document
from core.engine import FALLBACK_WARNING, dispatch


def make_minimal_docx(path: str) -> None:
    doc = Document()
    doc.add_paragraph("Engine Doc")
    doc.save(path)


def test_dispatch_headless_tools(tmp_path):
    f = tmp_path / "d.docx"
    make_minimal_docx(str(f))

    res = dispatch("read_text", {"path": str(f), "start": 0, "end": 0, "mode": "headless"})
    assert res["paragraphs"][0]["text"] == "Engine Doc"

    add_res = dispatch("add_paragraph", {"path": str(f), "text": "X", "mode": "headless"})
    assert add_res["ok"] is True


def test_dispatch_invalid_mode_raises(tmp_path):
    f = tmp_path / "d.docx"
    make_minimal_docx(str(f))
    with pytest.raises(ValueError, match="unknown mode"):
        dispatch("read_text", {"path": str(f), "start": 0, "end": 0, "mode": "invalid_mode"})


def test_dispatch_unknown_tool_raises(tmp_path):
    f = tmp_path / "d.docx"
    make_minimal_docx(str(f))
    with pytest.raises(ValueError, match="unknown tool"):
        dispatch("unknown_tool", {"path": str(f), "mode": "headless"})
