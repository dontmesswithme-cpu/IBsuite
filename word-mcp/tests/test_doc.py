import json
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from core.format_headless import apply_theme

FIXTURE_PATH = Path(__file__).parent / "fixtures" / "ib_memo.docx"
GOLDEN_PATH = Path(__file__).parent / "golden_doc.json"
GOLDEN = json.loads(GOLDEN_PATH.read_text(encoding="utf-8"))


def _rgb_hex(rgb_color) -> str | None:
    if rgb_color is None:
        return None
    try:
        return f"{rgb_color[0]:02X}{rgb_color[1]:02X}{rgb_color[2]:02X}"
    except Exception:
        return None


def test_golden_schema_frozen():
    assert set(GOLDEN.keys()) == {"fills", "fonts", "layout"}


def test_full_ib_grade_vs_golden():
    assert FIXTURE_PATH.is_file()
    doc = Document(str(FIXTURE_PATH))
    assert len(doc.sections) >= 1
    # Letter
    assert round(doc.sections[0].page_width.inches, 2) == GOLDEN["layout"]["page_width_inches"]
    assert round(doc.sections[0].page_height.inches, 2) == GOLDEN["layout"]["page_height_inches"]

    # Heading fill and font
    title = doc.paragraphs[0]
    shd = title._p.get_or_add_pPr().find(qn("w:shd"))
    assert shd is not None
    assert shd.get(qn("w:fill")).upper() == GOLDEN["fills"]["heading"]
    assert title.runs[0].font.name == GOLDEN["fonts"]["heading"]["name"]
    assert title.runs[0].font.bold is GOLDEN["fonts"]["heading"]["bold"]
    assert _rgb_hex(title.runs[0].font.color.rgb) == GOLDEN["fonts"]["heading"]["color"]

    # Table header fill and font
    assert len(doc.tables) >= 1
    t = doc.tables[0]
    header_cell = t.cell(0, 0)
    tc_shd = header_cell._tc.get_or_add_tcPr().find(qn("w:shd"))
    assert tc_shd is not None
    assert tc_shd.get(qn("w:fill")).upper() == GOLDEN["fills"]["heading"]
    cell_p = header_cell.paragraphs[0]
    assert cell_p.runs[0].font.bold is True
    assert _rgb_hex(cell_p.runs[0].font.color.rgb) == GOLDEN["fonts"]["heading"]["color"]


def test_apply_theme_matches_golden(tmp_path):
    doc_path = tmp_path / "test_memo.docx"
    doc = Document()
    p = doc.add_paragraph("Test Heading")
    p.style = "Heading 1"
    doc.add_paragraph("Body paragraph text")
    doc.save(str(doc_path))

    res = apply_theme(str(doc_path))
    assert res["ok"] is True

    doc2 = Document(str(doc_path))
    assert round(doc2.sections[0].page_width.inches, 2) == GOLDEN["layout"]["page_width_inches"]
    assert round(doc2.sections[0].page_height.inches, 2) == GOLDEN["layout"]["page_height_inches"]

    h = doc2.paragraphs[0]
    shd = h._p.get_or_add_pPr().find(qn("w:shd"))
    assert shd is not None
    assert shd.get(qn("w:fill")).upper() == GOLDEN["fills"]["heading"]
    assert h.runs[0].font.name == GOLDEN["fonts"]["heading"]["name"]
    assert h.runs[0].font.bold is GOLDEN["fonts"]["heading"]["bold"]
    assert _rgb_hex(h.runs[0].font.color.rgb) == GOLDEN["fonts"]["heading"]["color"]


def test_fixture_is_synthetic_only():
    doc = Document(str(FIXTURE_PATH))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Investment Banking Memo" in text
    assert "CONFIDENTIAL" not in text or "PROJECT" not in text
