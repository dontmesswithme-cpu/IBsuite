"""Task 1 tests: copy-before-write backups, FIFO retention, byte-identical rollback, audit log (word)."""

import hashlib
import json

from docx import Document

import core.audit as audit
import core.versions as versions
from core.audit import append
from core.versions import backup, list_versions, rollback

MARK = "orig"


def write_minimal_docx(path: str) -> None:
    doc = Document()
    doc.add_paragraph(MARK)
    doc.save(path)


def corrupt(path: str) -> None:
    doc = Document()
    doc.add_paragraph("CLOBBERED")
    doc.save(path)


def read_back(path: str) -> str:
    doc = Document(path)
    return doc.paragraphs[0].text


def sha(path: str) -> str:
    with open(path, "rb") as fh:
        return hashlib.sha256(fh.read()).hexdigest()


def paragraph_count(path: str) -> int:
    return len(Document(path).paragraphs)


def test_backup_restore_on_apply_crash(tmp_path):
    f = tmp_path / "d.docx"
    write_minimal_docx(str(f))
    original_sha = sha(str(f))

    bid = backup(str(f), "add_paragraph", "para-1")
    assert bid
    assert len(list_versions(str(f))) == 1

    corrupt(str(f))
    assert read_back(str(f)) != MARK

    result = rollback(str(f), bid)
    assert result["ok"] is True
    assert read_back(str(f)) == MARK
    assert sha(str(f)) == original_sha


def test_word_backup_fifo(tmp_path):
    f = tmp_path / "d.docx"
    write_minimal_docx(str(f))
    for i in range(21):
        backup(str(f), "t", f"r{i}")
    assert len(list_versions(str(f))) == 20


def test_fifo_evicts_oldest_first(tmp_path):
    f = tmp_path / "d.docx"
    write_minimal_docx(str(f))

    ids = [backup(str(f), "t", f"r{i}") for i in range(21)]
    retained = list_versions(str(f))

    assert ids[0] not in retained
    assert set(retained) == set(ids[1:])


def test_rollback_backs_up_pre_rollback_state(tmp_path):
    f = tmp_path / "d.docx"
    write_minimal_docx(str(f))

    bid = backup(str(f), "add_paragraph", "para-1")
    corrupt(str(f))

    rollback(str(f), bid)

    versions_for_file = list_versions(str(f))
    assert bid in versions_for_file
    assert len(versions_for_file) == 2


def test_backup_writes_json_sidecar(tmp_path):
    f = tmp_path / "d.docx"
    write_minimal_docx(str(f))

    bid = backup(str(f), "add_paragraph", "para-1")
    sidecar = tmp_path / "d.versions" / (bid[:-4] + ".json")

    assert sidecar.exists()
    meta = json.loads(sidecar.read_text(encoding="utf-8"))
    assert meta["tool"] == "add_paragraph"
    assert meta["label"] == "para-1"
    assert meta["backup_id"] == bid


def test_rollback_unknown_backup_id_rejected(tmp_path):
    f = tmp_path / "d.docx"
    write_minimal_docx(str(f))

    result = rollback(str(f), "not-a-real-backup.bak")
    assert result["ok"] is False


def test_versions_dir_is_stem_sibling(tmp_path):
    f = tmp_path / "memo.docx"
    write_minimal_docx(str(f))

    backup(str(f), "add_paragraph", "para-1")

    assert (tmp_path / "memo.versions").is_dir()
    assert (tmp_path / "memo.versions").parent == tmp_path


def test_minimal_docx_has_one_paragraph(tmp_path):
    f = tmp_path / "d.docx"
    write_minimal_docx(str(f))

    assert paragraph_count(str(f)) == 1
    assert read_back(str(f)) == MARK


def test_audit_append_writes_pipe_delimited_line(tmp_path, monkeypatch):
    log = tmp_path / "audit.log"
    monkeypatch.setattr(audit, "AUDIT_LOG_PATH", log)

    append("add_paragraph", "memo.docx", "para:1", "bid123", "ok")

    fields = [f.strip() for f in log.read_text(encoding="utf-8").strip().split("|")]
    assert fields[1] == "add_paragraph"
    assert fields[2] == "memo.docx"
    assert fields[3] == "para:1"
    assert fields[4] == "bid123"
    assert fields[5] == "ok"


def test_audit_append_creates_log_and_appends(tmp_path, monkeypatch):
    log = tmp_path / "audit.log"
    monkeypatch.setattr(audit, "AUDIT_LOG_PATH", log)

    append("preview_diff", "memo.docx", "diff", None, "ok")
    append("add_paragraph", "memo.docx", "para:2", "b2", "ERROR")

    lines = log.read_text(encoding="utf-8").strip().splitlines()
    assert len(lines) == 2
    assert lines[0].split("|")[4].strip() == "-"
    assert "ERROR" in lines[1]


def test_audit_line_is_append_only(tmp_path, monkeypatch):
    log = tmp_path / "audit.log"
    monkeypatch.setattr(audit, "AUDIT_LOG_PATH", log)

    for i in range(5):
        append("preview_diff", "memo.docx", f"diff:{i}", None, "ok")

    assert len(log.read_text(encoding="utf-8").strip().splitlines()) == 5


def test_audit_default_path_is_repo_audit_log():
    assert audit.AUDIT_LOG_PATH.name == "audit.log"
    assert "word-mcp" in str(audit.AUDIT_LOG_PATH)


def test_versions_retention_constant_is_twenty():
    assert versions.MAX_VERSIONS == 20
