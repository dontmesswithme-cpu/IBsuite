from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT


def _open_existing(path: str) -> Document:
    target = Path(path)
    if not target.is_file():
        raise FileNotFoundError(f"document not found: {target}")
    return Document(str(target))


def _save(doc: Document, path: str) -> None:
    doc.save(path)


def list_structure(path: str) -> list[dict]:
    doc = _open_existing(path)
    items = []
    for idx, para in enumerate(doc.paragraphs):
        items.append({"index": idx, "style": para.style.name, "text": para.text})
    return items


def read_text(path: str, start: int, end: int) -> dict:
    doc = _open_existing(path)
    paragraphs = doc.paragraphs
    total = len(paragraphs)
    if start < 0 or end < 0 or start >= total or end >= total or start > end:
        raise ValueError(f"range out of bounds: {start}-{end} (total {total})")
    selected = []
    for idx in range(start, end + 1):
        p = paragraphs[idx]
        selected.append({"index": idx, "style": p.style.name, "text": p.text})
    return {"paragraphs": selected, "total": total}


def add_paragraph(path: str, text: str = "", style: str | None = None) -> dict:
    doc = _open_existing(path)
    paragraph = doc.add_paragraph(text)
    if style:
        try:
            paragraph.style = style
        except KeyError:
            paragraph.style = "Normal"
    _save(doc, path)
    return {"ok": True, "index": len(doc.paragraphs) - 1, "text": text, "style": paragraph.style.name}


def set_heading(path: str, index: int, level: int, text: str) -> dict:
    doc = _open_existing(path)
    paragraphs = doc.paragraphs
    if index < 0:
        raise ValueError(f"unknown paragraph index: {index}")
    if index < len(paragraphs):
        para = paragraphs[index]
    else:
        # append empty paragraphs until index exists
        while len(doc.paragraphs) <= index:
            doc.add_paragraph("")
        para = doc.paragraphs[index]
    style_name = f"Heading {level}"
    try:
        para.style = style_name
    except KeyError:
        try:
            para.style = "Heading"
        except KeyError:
            pass
    para.text = text
    _save(doc, path)
    return {"ok": True, "index": index, "level": level, "text": text, "style": para.style.name}


def format_text(path: str, start: int, end: int, style: str) -> dict:
    doc = _open_existing(path)
    paragraphs = doc.paragraphs
    total = len(paragraphs)
    if start < 0 or end >= total or start > end:
        raise ValueError(f"range out of bounds: {start}-{end}")
    for idx in range(start, end + 1):
        try:
            paragraphs[idx].style = style
        except KeyError:
            pass
    _save(doc, path)
    return {"ok": True, "start": start, "end": end, "style": style}


def add_table(path: str, anchor, rows: int, cols: int, values=None) -> dict:
    doc = _open_existing(path)
    if anchor == "end" or anchor is None:
        table = doc.add_table(rows=rows, cols=cols)
        table.style = "Table Grid"
    else:
        # anchor as paragraph index: insert after that paragraph
        try:
            idx = int(anchor)
        except (TypeError, ValueError):
            idx = None
        if idx is None:
            table = doc.add_table(rows=rows, cols=cols)
            table.style = "Table Grid"
        else:
            # python-docx doesn't support inserting table at arbitrary position easily;
            # append at end for simplicity
            table = doc.add_table(rows=rows, cols=cols)
            table.style = "Table Grid"
    if values:
        for r_idx, row in enumerate(values):
            if r_idx >= rows:
                break
            for c_idx, val in enumerate(row):
                if c_idx >= cols:
                    break
                table.cell(r_idx, c_idx).text = str(val)
    _save(doc, path)
    return {"ok": True, "rows": rows, "cols": cols, "anchor": anchor}


def add_image(path: str, anchor, image_path: str) -> dict:
    from pathlib import Path as _P

    img = _P(image_path)
    if not img.is_file():
        raise FileNotFoundError(f"image not found: {image_path}")
    doc = _open_existing(path)
    doc.add_picture(str(img))
    _save(doc, path)
    return {"ok": True, "image_path": str(image_path), "anchor": anchor}


def add_header_footer(path: str, kind: str, text: str) -> dict:
    doc = _open_existing(path)
    section = doc.sections[0]
    if kind == "header":
        section.header.paragraphs[0].text = text
    elif kind == "footer":
        section.footer.paragraphs[0].text = text
    else:
        raise ValueError(f"unknown kind: {kind!r}")
    _save(doc, path)
    return {"ok": True, "kind": kind, "text": text}


def _shade_paragraph(paragraph, hex_color: str) -> None:
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def apply_theme(path: str) -> dict:
    doc = _open_existing(path)
    from docx.shared import Pt, RGBColor

    # Normal style
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    header_fill = "002060"
    for para in doc.paragraphs:
        if para.style.name.startswith("Heading"):
            _shade_paragraph(para, header_fill)
            if not para.runs and para.text:
                para.add_run(para.text)
            for run in para.runs:
                run.font.name = "Arial"
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    # tables
    for table in doc.tables:
        if table.rows:
            for cell in table.rows[0].cells:
                _shade_cell(cell, header_fill)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.name = "Arial"
            for row in table.rows[1:]:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.name = "Arial"

    # page setup Letter
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)

    _save(doc, path)
    return {"ok": True, "applied": "ib_doc_theme"}


def _shade_cell(cell, hex_color: str) -> None:
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_page_setup(path: str, size: str = "Letter", margins=None, orientation: str = "portrait") -> dict:
    doc = _open_existing(path)
    sizes = {
        "Letter": (Inches(8.5), Inches(11)),
        "A4": (Inches(8.27), Inches(11.69)),
    }
    w, h = sizes.get(size, sizes["Letter"])
    for section in doc.sections:
        if orientation == "landscape":
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width = h
            section.page_height = w
        else:
            section.orientation = WD_ORIENT.PORTRAIT
            section.page_width = w
            section.page_height = h
        if margins is not None:
            try:
                m = float(margins)
                section.top_margin = Inches(m)
                section.bottom_margin = Inches(m)
                section.left_margin = Inches(m)
                section.right_margin = Inches(m)
            except (TypeError, ValueError):
                if isinstance(margins, dict):
                    if "top" in margins:
                        section.top_margin = Inches(float(margins["top"]))
                    if "bottom" in margins:
                        section.bottom_margin = Inches(float(margins["bottom"]))
                    if "left" in margins:
                        section.left_margin = Inches(float(margins["left"]))
                    if "right" in margins:
                        section.right_margin = Inches(float(margins["right"]))
    _save(doc, path)
    return {"ok": True, "size": size, "orientation": orientation}


def preview_diff(path: str, ops: dict) -> dict:
    doc = _open_existing(path)
    return {
        "ok": True,
        "diff": {
            "path": path,
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
            "proposed_operation": ops,
        },
    }

