import base64
import io
import os
from pathlib import Path
import tempfile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

THUMB_WIDTH = 320
THUMB_MAX_HEIGHT = 240
FULL_WIDTH = 1280
MAX_THUMB_B64 = 40000


def _load_font(size: int = 16, bold: bool = False):
    candidates = []
    if bold:
        candidates.append(r"C:\Windows\Fonts\arialbd.ttf")
    candidates.append(r"C:\Windows\Fonts\arial.ttf")
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            continue
    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()


def _render_headless(path: str, slide: int = 1) -> Image.Image:
    target = Path(path).resolve()
    if not target.is_file():
        raise FileNotFoundError(f"presentation not found: {target}")

    prs = Presentation(str(target))
    if slide < 1 or slide > len(prs.slides):
        raise ValueError(f"unknown slide: {slide}")

    sld = prs.slides[slide - 1]

    # Standard 16:9 canvas
    width, height = 1280, 720
    img = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(img)

    title_text = ""
    if sld.shapes.title and sld.shapes.title.has_text_frame:
        title_text = sld.shapes.title.text

    # Draw header bar if title exists
    if title_text:
        draw.rectangle([0, 0, width, 100], fill="#002060")
        font = _load_font(28, bold=True)
        draw.text((40, 30), title_text, font=font, fill="#FFFFFF")

    # Draw shapes / text / table outlines
    y_offset = 140
    for shape in sld.shapes:
        if shape == sld.shapes.title:
            continue
        if shape.has_text_frame and shape.text_frame.text.strip():
            font = _load_font(18, bold=False)
            text_lines = shape.text_frame.text.splitlines()[:5]
            for line in text_lines:
                draw.text((40, y_offset), line[:80], font=font, fill="#333333")
                y_offset += 26
        elif shape.has_table:
            draw.rectangle([40, y_offset, width - 40, y_offset + 180], outline="#D9D9D9", width=2)
            draw.rectangle([40, y_offset, width - 40, y_offset + 35], fill="#002060")
            font = _load_font(14, bold=True)
            draw.text((50, y_offset + 8), "Table Content", font=font, fill="#FFFFFF")
            y_offset += 200

    return img


def _render_headed(path: str, slide: int = 1) -> Image.Image | None:
    try:
        import win32com.client
    except ImportError:
        return None

    target = Path(path).resolve()
    if not target.is_file():
        return None

    try:
        app = win32com.client.Dispatch("PowerPoint.Application")
    except Exception:
        return None

    try:
        pres = app.Presentations.Open(str(target), ReadOnly=True, WithWindow=False)
        try:
            if slide < 1 or slide > pres.Slides.Count:
                return None
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                tmp_path = tmp.name
            try:
                pres.Slides(slide).Export(tmp_path, "PNG", FULL_WIDTH, int(FULL_WIDTH * 9 / 16))
                if os.path.exists(tmp_path) and os.path.getsize(tmp_path) > 0:
                    with Image.open(tmp_path) as exported:
                        return exported.convert("RGB")
            finally:
                if os.path.exists(tmp_path):
                    os.unlink(tmp_path)
        finally:
            try:
                pres.Close()
            except Exception:
                pass
    except Exception:
        return None
    finally:
        try:
            app.Quit()
        except Exception:
            pass
    return None


def _to_b64(img: Image.Image) -> str:
    buffer = io.BytesIO()
    img.save(buffer, format="PNG", optimize=True)
    return base64.b64encode(buffer.getvalue()).decode("ascii")


def _fit(img: Image.Image, width: int) -> Image.Image:
    current_w, current_h = img.size
    if current_w == width:
        return img
    new_h = max(1, int(current_h * width / current_w))
    return img.resize((width, new_h), Image.LANCZOS)


def thumbnail(path: str, slide: int = 1) -> str:
    img = _render_headless(path, slide)
    small = _fit(img, THUMB_WIDTH)
    if small.size[1] > THUMB_MAX_HEIGHT:
        small = small.crop((0, 0, small.size[0], THUMB_MAX_HEIGHT))
    small = small.convert("P", palette=Image.ADAPTIVE, colors=32).convert("RGB")
    out = _to_b64(small)
    while len(out) >= MAX_THUMB_B64 and small.size[0] > 160:
        small = small.resize((small.size[0] // 2, small.size[1] // 2), Image.LANCZOS)
        small = small.convert("P", palette=Image.ADAPTIVE, colors=16).convert("RGB")
        out = _to_b64(small)
    return out


def full_shot(path: str, slide: int = 1, mode: str = "headless") -> str:
    if mode not in ("headless", "headed"):
        raise ValueError(f"unknown mode: {mode!r} (expected 'headless' or 'headed')")
    if mode == "headed":
        shot = _render_headed(path, slide)
        if shot is not None:
            return _to_b64(_fit(shot, FULL_WIDTH))
    return _to_b64(_fit(_render_headless(path, slide), FULL_WIDTH))
