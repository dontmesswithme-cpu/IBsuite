from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from core.format_headless import apply_theme, format_shape


def _open_existing(path: str) -> Presentation:
    target = Path(path)
    if not target.is_file():
        raise FileNotFoundError(f"presentation not found: {target}")
    return Presentation(str(target))


def list_slides(path: str) -> list[dict]:
    prs = _open_existing(path)
    slides = []
    for idx, slide in enumerate(prs.slides, start=1):
        title = ""
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text
        layout_name = slide.slide_layout.name if slide.slide_layout else "unknown"
        slides.append(
            {
                "slide": idx,
                "title": title,
                "layout": layout_name,
                "shapes_count": len(slide.shapes),
            }
        )
    return slides


def read_text(path: str, slide: int) -> dict:
    prs = _open_existing(path)
    if slide < 1 or slide > len(prs.slides):
        raise ValueError(f"unknown slide: {slide}")

    slide_obj = prs.slides[slide - 1]
    title = ""
    if slide_obj.shapes.title and slide_obj.shapes.title.has_text_frame:
        title = slide_obj.shapes.title.text
    else:
        for s in slide_obj.shapes:
            if s.has_text_frame and "title" in s.name.lower():
                title = s.text_frame.text
                break

    shapes_text = {}
    all_text = []

    for shape in slide_obj.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text
            shapes_text[shape.name] = txt
            if txt.strip():
                all_text.append(txt)
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    txt = cell.text_frame.text
                    if txt.strip():
                        all_text.append(txt)

    return {
        "slide": slide,
        "title": title,
        "shapes": shapes_text,
        "text": "\n".join(all_text),
    }


def add_slide(path: str, layout: str = "blank", index: int | None = None) -> dict:
    prs = _open_existing(path)

    layout_map = {
        "title": 0,
        "title_content": 1,
        "content": 1,
        "section": 2,
        "two_content": 3,
        "comparison": 4,
        "title_only": 5,
        "blank": 6,
    }

    chosen_idx = layout_map.get(str(layout).lower(), 6)
    if chosen_idx >= len(prs.slide_layouts):
        chosen_idx = 0
    slide_layout = prs.slide_layouts[chosen_idx]

    prs.slides.add_slide(slide_layout)
    prs.save(path)

    total = len(prs.slides)
    return {"ok": True, "slide": total, "total_slides": total}


def set_text(path: str, slide: int, shape: str, text: str) -> dict:
    prs = _open_existing(path)
    if slide < 1 or slide > len(prs.slides):
        raise ValueError(f"unknown slide: {slide}")

    slide_obj = prs.slides[slide - 1]
    target_shape = None

    if shape.lower() == "title":
        target_shape = slide_obj.shapes.title
        if not target_shape:
            target_shape = next(
                (s for s in slide_obj.shapes if s.has_text_frame and "title" in s.name.lower()),
                None,
            )
        if not target_shape:
            target_shape = slide_obj.shapes.add_textbox(
                Inches(1), Inches(0.5), Inches(11.33), Inches(1.2)
            )
            target_shape.name = "Title"
    else:
        target_shape = next((s for s in slide_obj.shapes if s.name == shape), None)

    if not target_shape:
        raise ValueError(f"unknown shape: {shape!r}")

    if not target_shape.has_text_frame:
        raise ValueError(f"shape {shape!r} does not support text")

    target_shape.text_frame.text = text
    prs.save(path)

    return {"ok": True, "slide": slide, "shape": shape, "text": text}


def add_table(
    path: str,
    slide: int,
    rows: int,
    cols: int,
    values: list[list] | None = None,
    anchor: str | tuple | None = None,
) -> dict:
    prs = _open_existing(path)
    if slide < 1 or slide > len(prs.slides):
        raise ValueError(f"unknown slide: {slide}")

    slide_obj = prs.slides[slide - 1]
    left = Inches(1)
    top = Inches(2)
    width = Inches(11.33)
    height = Inches(4)

    table_shape = slide_obj.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    if values:
        for r_idx, row in enumerate(values):
            if r_idx >= rows:
                break
            for c_idx, val in enumerate(row):
                if c_idx >= cols:
                    break
                table.cell(r_idx, c_idx).text = str(val)

    prs.save(path)
    return {"ok": True, "slide": slide, "rows": rows, "cols": cols}


def add_chart(
    path: str,
    slide: int,
    chart_type: str,
    data: dict,
    anchor: str | tuple | None = None,
) -> dict:
    prs = _open_existing(path)
    if slide < 1 or slide > len(prs.slides):
        raise ValueError(f"unknown slide: {slide}")

    slide_obj = prs.slides[slide - 1]

    chart_type_map = {
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "line": XL_CHART_TYPE.LINE,
        "pie": XL_CHART_TYPE.PIE,
    }
    xl_type = chart_type_map.get(str(chart_type).lower(), XL_CHART_TYPE.COLUMN_CLUSTERED)

    chart_data = CategoryChartData()
    chart_data.categories = data.get("cats", ["Cat 1", "Cat 2"])
    series_dict = data.get("series", {"Series 1": [1, 2]})
    for name, vals in series_dict.items():
        chart_data.add_series(name, tuple(vals))

    left = Inches(1)
    top = Inches(1.5)
    width = Inches(11.33)
    height = Inches(5)

    slide_obj.shapes.add_chart(xl_type, left, top, width, height, chart_data)
    prs.save(path)

    return {"ok": True, "slide": slide, "chart_type": chart_type}


def add_image(
    path: str,
    slide: int,
    image_path: str,
    anchor: str | tuple | None = None,
) -> dict:
    prs = _open_existing(path)
    if slide < 1 or slide > len(prs.slides):
        raise ValueError(f"unknown slide: {slide}")

    img_file = Path(image_path)
    if not img_file.is_file():
        raise FileNotFoundError(f"image not found: {image_path}")

    slide_obj = prs.slides[slide - 1]
    left = Inches(1)
    top = Inches(1.5)

    slide_obj.shapes.add_picture(str(img_file), left, top)
    prs.save(path)

    return {"ok": True, "slide": slide, "image_path": str(image_path)}


def preview_diff(path: str, ops: dict) -> dict:
    prs = _open_existing(path)
    slide_count = len(prs.slides)
    return {
        "ok": True,
        "diff": {
            "path": path,
            "current_slide_count": slide_count,
            "proposed_operation": ops,
        },
    }
