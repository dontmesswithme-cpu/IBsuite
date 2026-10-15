from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

import config


def _hex_to_rgb(hex_code: str) -> RGBColor:
    clean = hex_code.strip().lstrip("#")
    if len(clean) == 3:
        clean = "".join(ch * 2 for ch in clean)
    if len(clean) != 6:
        clean = "000000"
    r = int(clean[0:2], 16)
    g = int(clean[2:4], 16)
    b = int(clean[4:6], 16)
    return RGBColor(r, g, b)


def _open_existing(path: str) -> Presentation:
    target = Path(path)
    if not target.is_file():
        raise FileNotFoundError(f"presentation not found: {target}")
    return Presentation(str(target))


def format_shape(
    path: str,
    slide: int,
    shape: str,
    font_name: str | None = None,
    font_size: int | None = None,
    font_color: str | None = None,
    fill_color: str | None = None,
    bold: bool | None = None,
) -> dict:
    prs = _open_existing(path)
    if slide < 1 or slide > len(prs.slides):
        raise ValueError(f"unknown slide: {slide}")

    slide_obj = prs.slides[slide - 1]
    target_shape = None

    if shape.lower() == "title":
        target_shape = slide_obj.shapes.title
        if not target_shape:
            target_shape = next(
                (s for s in slide_obj.shapes if s.has_text_frame and "title" in s.name.lower()),
                None,
            )
    else:
        target_shape = next((s for s in slide_obj.shapes if s.name == shape), None)

    if not target_shape:
        raise ValueError(f"unknown shape: {shape!r}")

    if fill_color is not None:
        target_shape.fill.solid()
        target_shape.fill.fore_color.rgb = _hex_to_rgb(fill_color)

    if target_shape.has_text_frame:
        for paragraph in target_shape.text_frame.paragraphs:
            if font_name is not None:
                paragraph.font.name = font_name
            if font_size is not None:
                paragraph.font.size = Pt(font_size)
            if font_color is not None:
                paragraph.font.color.rgb = _hex_to_rgb(font_color)
            if bold is not None:
                paragraph.font.bold = bold

    prs.save(path)
    return {"ok": True, "slide": slide, "shape": shape}


def apply_theme(path: str) -> dict:
    prs = _open_existing(path)

    header_fill = str(config.resolve_style("header_fill", None))
    header_font_color = str(config.resolve_style("header_font_color", None))
    header_font_bold = bool(config.resolve_style("header_font_bold", None))
    body_font_name = str(config.resolve_style("body_font_name", None))
    table_grid_color = str(config.resolve_style("table_grid_color", None))

    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for slide in prs.slides:
        title = slide.shapes.title
        if title:
            title.fill.solid()
            title.fill.fore_color.rgb = _hex_to_rgb(header_fill)
            if title.has_text_frame:
                for p in title.text_frame.paragraphs:
                    p.font.name = body_font_name
                    p.font.bold = header_font_bold
                    p.font.color.rgb = _hex_to_rgb(header_font_color)

        for shape in slide.shapes:
            if shape.has_table:
                table = shape.table
                for col_idx in range(len(table.columns)):
                    cell = table.cell(0, col_idx)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _hex_to_rgb(header_fill)
                    for p in cell.text_frame.paragraphs:
                        p.font.name = body_font_name
                        p.font.bold = True
                        p.font.color.rgb = _hex_to_rgb(header_font_color)

                for row_idx in range(1, len(table.rows)):
                    for col_idx in range(len(table.columns)):
                        cell = table.cell(row_idx, col_idx)
                        for p in cell.text_frame.paragraphs:
                            p.font.name = body_font_name

    prs.save(path)
    return {"ok": True, "applied": "ib_deck_theme"}
