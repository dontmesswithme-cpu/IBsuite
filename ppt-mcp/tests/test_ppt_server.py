import pytest
from pptx import Presentation

import config
import core.audit as audit_mod
from server import call_tool, ops_hash


def make_minimal_pptx(path: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Server Test"
    prs.save(path)


@pytest.fixture
def sandbox(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "ALLOWED_ROOTS", [str(tmp_path)])
    log = tmp_path / "audit.log"
    monkeypatch.setattr(audit_mod, "AUDIT_LOG_PATH", log)
    return tmp_path, log


def test_read_tools_need_no_approval(sandbox):
    tmp_path, _log = sandbox
    f = tmp_path / "d.pptx"
    make_minimal_pptx(str(f))

    assert call_tool("list_slides", {"path": str(f)})["ok"] is True
    out = call_tool("read_text", {"path": str(f), "slide": 1})
    assert out["ok"] is True and out["title"] == "Server Test"

    diff = call_tool(
        "preview_diff",
        {"path": str(f), "ops": {"tool": "add_slide", "layout": "blank"}},
    )
    assert diff["ok"] is True


def test_write_needs_approval_then_succeeds(sandbox):
    tmp_path, _log = sandbox
    f = tmp_path / "d.pptx"
    make_minimal_pptx(str(f))

    ops = {
        "path": str(f),
        "slide": 1,
        "shape": "title",
        "text": "Approved Title",
    }
    first = call_tool("set_text", ops)
    assert first["status"] == "needs_approval"
    assert first["allowed"] is False

    ops["approval_token"] = first["ops_hash"]
    second = call_tool("set_text", ops)
    assert second["ok"] is True and second["backup_id"]

    prs = Presentation(str(f))
    assert prs.slides[0].shapes.title.text == "Approved Title"


def test_outside_path_rejected_before_approval(sandbox):
    tmp_path, log = sandbox
    outside = "C:\\Windows\\secret.pptx"

    denied = call_tool("read_text", {"path": outside, "slide": 1})
    assert denied["status"] == "rejected" and denied["ok"] is False

    assert not log.exists()


def test_rollback_needs_approval_then_restores(sandbox):
    tmp_path, _log = sandbox
    f = tmp_path / "d.pptx"
    make_minimal_pptx(str(f))

    write = {"path": str(f), "slide": 1, "shape": "title", "text": "New Text"}
    write["approval_token"] = ops_hash("set_text", write)
    first = call_tool("set_text", write)
    assert first["ok"] is True

    roll = {"path": str(f), "backup_id": first["backup_id"]}
    assert call_tool("rollback", roll)["status"] == "needs_approval"

    roll["approval_token"] = ops_hash("rollback", roll)
    done = call_tool("rollback", roll)
    assert done["ok"] is True and done["backup_id"]

    prs = Presentation(str(f))
    assert prs.slides[0].shapes.title.text == "Server Test"
