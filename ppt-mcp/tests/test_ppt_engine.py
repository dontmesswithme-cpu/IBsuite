import pytest
from pptx import Presentation

from core.engine import FALLBACK_WARNING, dispatch


def make_minimal_pptx(path: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Engine Slide"
    prs.save(path)


def test_dispatch_headless_tools(tmp_path):
    f = tmp_path / "deck.pptx"
    make_minimal_pptx(str(f))

    res = dispatch("read_text", {"path": str(f), "slide": 1, "mode": "headless"})
    assert res["title"] == "Engine Slide"

    add_res = dispatch("add_slide", {"path": str(f), "layout": "blank", "mode": "headless"})
    assert add_res["ok"] is True
    assert add_res["slide"] == 2


def test_dispatch_invalid_mode_raises(tmp_path):
    f = tmp_path / "deck.pptx"
    make_minimal_pptx(str(f))

    with pytest.raises(ValueError, match="unknown mode"):
        dispatch("read_text", {"path": str(f), "slide": 1, "mode": "invalid_mode"})


def test_dispatch_unknown_tool_raises(tmp_path):
    f = tmp_path / "deck.pptx"
    make_minimal_pptx(str(f))

    with pytest.raises(ValueError, match="unknown tool"):
        dispatch("unknown_tool", {"path": str(f), "mode": "headless"})
