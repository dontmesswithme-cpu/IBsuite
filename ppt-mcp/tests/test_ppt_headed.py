import pytest
from pptx import Presentation

from core.engine import dispatch
from core.headed import HeadedUnavailable, _powerpoint, is_powerpoint_available


def make_minimal_pptx(path: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Sample Title"
    prs.save(path)


def test_headed_visible_default(monkeypatch):
    import core.headed as headed_mod

    seen = {}

    class FakeApp:
        def __init__(self, seen):
            self._seen = seen

        def __setattr__(self, name, value):
            if name in ("Visible", "DisplayAlerts"):
                self._seen[name] = value
            else:
                object.__setattr__(self, name, value)

        def Quit(self):
            pass

    monkeypatch.setattr(headed_mod, "_launch", lambda: FakeApp(seen))
    with headed_mod._powerpoint():
        pass
    assert seen["Visible"] is True


@pytest.mark.needs_powerpoint
def test_is_powerpoint_available_returns_bool():
    assert isinstance(is_powerpoint_available(), bool)


@pytest.mark.needs_powerpoint
def test_headed_unavailable_falls_back(tmp_path):
    f = tmp_path / "d.pptx"
    make_minimal_pptx(str(f))
    out = dispatch("read_text", {"path": str(f), "slide": 1, "mode": "headed"})
    assert "title" in out or out.get("warning") == "headed_unavailable_fell_back"
