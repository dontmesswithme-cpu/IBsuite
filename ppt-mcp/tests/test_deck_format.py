from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import pytest

from core.format_headless import apply_theme, format_shape


def make_deck(path: str, slides: int = 1) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for i in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        if slide.shapes.title:
            slide.shapes.title.text = f"Slide {i+1} Title"
    prs.save(path)


def test_format_shape_custom_styling(tmp_path):
    f = tmp_path / "deck.pptx"
    make_deck(str(f), slides=1)

    res = format_shape(str(f), 1, "title", fill_color="002060", font_color="FFFFFF", bold=True, font_size=20)
    assert res["ok"] is True

    prs = Presentation(str(f))
    title = prs.slides[0].shapes.title
    assert title.fill.type is not None
    assert str(title.fill.fore_color.rgb).upper() == "002060"
    p = title.text_frame.paragraphs[0]
    assert p.font.bold is True
    assert str(p.font.color.rgb).upper() == "FFFFFF"


def test_theme_navy_golden(tmp_path):
    f = tmp_path / "d.pptx"
    make_deck(str(f), slides=1)

    res = apply_theme(str(f))
    assert res["ok"] is True

    prs = Presentation(str(f))
    slide = prs.slides[0]
    title = slide.shapes.title
    assert title is not None
    assert str(title.fill.fore_color.rgb).upper() == "002060"
    p = title.text_frame.paragraphs[0]
    assert p.font.bold is True
    assert str(p.font.color.rgb).upper() == "FFFFFF"
