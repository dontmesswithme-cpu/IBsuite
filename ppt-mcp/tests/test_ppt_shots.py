import base64
from pptx import Presentation
from pptx.util import Inches

from core.shots import full_shot, thumbnail


def make_deck(path: str) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Snapshot Test Slide"
    prs.save(path)


def test_thumbnail_returns_b64_under_40k(tmp_path):
    f = tmp_path / "deck.pptx"
    make_deck(str(f))

    thumb = thumbnail(str(f), slide=1)
    assert isinstance(thumb, str)
    assert len(thumb) < 40000
    decoded = base64.b64decode(thumb)
    assert len(decoded) > 0


def test_full_shot_larger_than_thumbnail(tmp_path):
    f = tmp_path / "deck.pptx"
    make_deck(str(f))

    thumb = thumbnail(str(f), slide=1)
    full = full_shot(str(f), slide=1, mode="headless")

    assert isinstance(full, str)
    assert len(full) > len(thumb)
