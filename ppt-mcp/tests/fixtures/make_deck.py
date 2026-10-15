from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt


def _hex_to_rgb(hex_code: str) -> RGBColor:
    clean = hex_code.strip().lstrip("#")
    return RGBColor(int(clean[0:2], 16), int(clean[2:4], 16), int(clean[4:6], 16))


def build(path: str) -> None:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    header_rgb = _hex_to_rgb("002060")
    white_rgb = _hex_to_rgb("FFFFFF")

    # Slide 1: Title Slide
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Investment Banking Pitch Book"
    s1.shapes.title.fill.solid()
    s1.shapes.title.fill.fore_color.rgb = header_rgb
    for p_elem in s1.shapes.title.text_frame.paragraphs:
        p_elem.font.name = "Arial"
        p_elem.font.bold = True
        p_elem.font.color.rgb = white_rgb
    if len(s1.placeholders) > 1:
        s1.placeholders[1].text = "M&A Advisory & Valuation Analysis"

    # Slide 2: Three-Statement Summary table
    s2 = prs.slides.add_slide(prs.slide_layouts[5])  # Title only
    s2.shapes.title.text = "Three-Statement Summary"
    s2.shapes.title.fill.solid()
    s2.shapes.title.fill.fore_color.rgb = header_rgb
    for p_elem in s2.shapes.title.text_frame.paragraphs:
        p_elem.font.name = "Arial"
        p_elem.font.bold = True
        p_elem.font.color.rgb = white_rgb

    t2_shape = s2.shapes.add_table(4, 4, Inches(1), Inches(2), Inches(11.333), Inches(3.5))
    t2 = t2_shape.table
    headers = ["Metric ($M)", "2024A", "2025E", "2026E"]
    for c_idx, h in enumerate(headers):
        cell = t2.cell(0, c_idx)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_rgb
        for p_elem in cell.text_frame.paragraphs:
            p_elem.font.name = "Arial"
            p_elem.font.bold = True
            p_elem.font.color.rgb = white_rgb

    rows = [
        ["Revenue", "1,250", "1,450", "1,700"],
        ["EBITDA", "310", "380", "460"],
        ["Net Income", "180", "225", "280"],
    ]
    for r_idx, r_data in enumerate(rows, start=1):
        for c_idx, val in enumerate(r_data):
            cell = t2.cell(r_idx, c_idx)
            cell.text = val
            for p_elem in cell.text_frame.paragraphs:
                p_elem.font.name = "Arial"

    # Slide 3: Public Comps Valuation table
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Public Comps Valuation"
    s3.shapes.title.fill.solid()
    s3.shapes.title.fill.fore_color.rgb = header_rgb
    for p_elem in s3.shapes.title.text_frame.paragraphs:
        p_elem.font.name = "Arial"
        p_elem.font.bold = True
        p_elem.font.color.rgb = white_rgb

    t3_shape = s3.shapes.add_table(4, 4, Inches(1), Inches(2), Inches(11.333), Inches(3.5))
    t3 = t3_shape.table
    c_headers = ["Peer Company", "EV/EBITDA", "P/E", "Mkt Cap ($B)"]
    for c_idx, h in enumerate(c_headers):
        cell = t3.cell(0, c_idx)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_rgb
        for p_elem in cell.text_frame.paragraphs:
            p_elem.font.name = "Arial"
            p_elem.font.bold = True
            p_elem.font.color.rgb = white_rgb

    c_rows = [
        ["Peer Corp A", "14.2x", "22.5x", "45.2"],
        ["Peer Corp B", "12.8x", "19.8x", "38.6"],
        ["Mean / Median", "13.5x", "21.1x", "41.9"],
    ]
    for r_idx, r_data in enumerate(c_rows, start=1):
        for c_idx, val in enumerate(r_data):
            cell = t3.cell(r_idx, c_idx)
            cell.text = val
            for p_elem in cell.text_frame.paragraphs:
                p_elem.font.name = "Arial"

    # Slide 4: Chart slide
    s4 = prs.slides.add_slide(prs.slide_layouts[5])
    s4.shapes.title.text = "Revenue & EBITDA Trajectory"
    s4.shapes.title.fill.solid()
    s4.shapes.title.fill.fore_color.rgb = header_rgb
    for p_elem in s4.shapes.title.text_frame.paragraphs:
        p_elem.font.name = "Arial"
        p_elem.font.bold = True
        p_elem.font.color.rgb = white_rgb

    chart_data = CategoryChartData()
    chart_data.categories = ["2024A", "2025E", "2026E"]
    chart_data.add_series("Revenue", (1250, 1450, 1700))
    chart_data.add_series("EBITDA", (310, 380, 460))
    s4.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(2),
        Inches(11.333),
        Inches(4.5),
        chart_data,
    )

    prs.save(str(p))


if __name__ == "__main__":
    out_path = Path(__file__).parent / "ib_deck.pptx"
    build(str(out_path))
    print(f"Generated {out_path}")
