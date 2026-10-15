import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

from core.format_headless import apply_theme

FIXTURE_PATH = Path(__file__).parent / "fixtures" / "ib_deck.pptx"
GOLDEN_PATH = Path(__file__).parent / "golden_deck.json"
GOLDEN = json.loads(GOLDEN_PATH.read_text(encoding="utf-8"))


def _rgb(color_format) -> str | None:
    try:
        return str(color_format.rgb).upper()
    except Exception:
        return None


def test_golden_schema_frozen():
    assert set(GOLDEN.keys()) == {"fills", "fonts", "layout", "slides"}


def test_full_ib_grade_vs_golden(tmp_path):
    assert FIXTURE_PATH.is_file()
    prs = Presentation(str(FIXTURE_PATH))

    # Assert layout & dimensions
    assert len(prs.slides) == GOLDEN["layout"]["slide_count"]
    assert round(prs.slide_width.inches, 3) == GOLDEN["layout"]["width_inches"]
    assert round(prs.slide_height.inches, 3) == GOLDEN["layout"]["height_inches"]

    for idx, expected in enumerate(GOLDEN["slides"]):
        slide = prs.slides[idx]
        title = slide.shapes.title
        assert title is not None
        assert title.text == expected["title"]

        # Check title formatting
        assert _rgb(title.fill.fore_color) == GOLDEN["fills"]["header"]
        p = title.text_frame.paragraphs[0]
        assert p.font.bold is GOLDEN["fonts"]["header"]["bold"]
        assert _rgb(p.font.color) == GOLDEN["fonts"]["header"]["color"]
        assert p.font.name == GOLDEN["fonts"]["header"]["name"]

        has_table = any(s.has_table for s in slide.shapes)
        has_chart = any(s.has_chart for s in slide.shapes)
        assert has_table is expected["has_table"]
        assert has_chart is expected["has_chart"]

        if has_table:
            table_shape = next(s for s in slide.shapes if s.has_table)
            t = table_shape.table
            header_cell = t.cell(0, 0)
            assert _rgb(header_cell.fill.fore_color) == GOLDEN["fills"]["header"]
            p_cell = header_cell.text_frame.paragraphs[0]
            assert p_cell.font.bold is True
            assert _rgb(p_cell.font.color) == GOLDEN["fonts"]["header"]["color"]


def test_apply_theme_matches_golden(tmp_path):
    deck_copy = tmp_path / "deck.pptx"
    deck_copy.write_bytes(FIXTURE_PATH.read_bytes())

    res = apply_theme(str(deck_copy))
    assert res["ok"] is True

    prs = Presentation(str(deck_copy))
    assert round(prs.slide_width.inches, 3) == GOLDEN["layout"]["width_inches"]
    assert round(prs.slide_height.inches, 3) == GOLDEN["layout"]["height_inches"]

    for slide in prs.slides:
        if slide.shapes.title:
            assert _rgb(slide.shapes.title.fill.fore_color) == GOLDEN["fills"]["header"]
            p = slide.shapes.title.text_frame.paragraphs[0]
            assert p.font.bold is True
            assert _rgb(p.font.color) == GOLDEN["fonts"]["header"]["color"]
