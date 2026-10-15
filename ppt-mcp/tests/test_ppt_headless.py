import hashlib
from pathlib import Path
from PIL import Image
import pytest
from pptx import Presentation
from pptx.util import Inches

from core.headless import (
    add_chart,
    add_image,
    add_slide,
    add_table,
    list_slides,
    preview_diff,
    read_text,
    set_text,
)


def make_deck(path: str, slides: int = 1) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for i in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        if slide.shapes.title:
            slide.shapes.title.text = f"Slide {i+1} Title"
    prs.save(path)


def make_test_image(path: str) -> None:
    img = Image.new("RGB", (100, 100), color="blue")
    img.save(path, format="PNG")


def test_list_slides(tmp_path):
    f = tmp_path / "deck.pptx"
    make_deck(str(f), slides=2)

    slides = list_slides(str(f))
    assert len(slides) == 2
    assert slides[0]["slide"] == 1
    assert slides[1]["slide"] == 2


def test_read_text_and_unknown_slide_raises(tmp_path):
    f = tmp_path / "deck.pptx"
    make_deck(str(f), slides=1)

    data = read_text(str(f), 1)
    assert data["title"] == "Slide 1 Title"

    with pytest.raises(ValueError, match="unknown slide"):
        read_text(str(f), 99)


def test_slide_text_table_roundtrip(tmp_path):
    f = tmp_path / "d.pptx"
    make_deck(str(f), slides=2)

    add_res = add_slide(str(f), "blank", None)
    assert add_res["ok"] is True
    assert add_res["slide"] == 3

    set_res = set_text(str(f), 3, "title", "Comps")
    assert set_res["ok"] is True
    assert read_text(str(f), 3)["title"] == "Comps"


def test_set_text_unknown_shape_raises(tmp_path):
    f = tmp_path / "d.pptx"
    make_deck(str(f), slides=1)

    with pytest.raises(ValueError, match="unknown shape"):
        set_text(str(f), 1, "nonexistent_shape_name_xyz", "Hello")


def test_add_table_and_read_values(tmp_path):
    f = tmp_path / "d.pptx"
    make_deck(str(f), slides=1)

    rows = 2
    cols = 2
    values = [["Metric", "Value"], ["Revenue", "$100M"]]
    res = add_table(str(f), 1, rows=rows, cols=cols, values=values)
    assert res["ok"] is True

    prs = Presentation(str(f))
    shapes = prs.slides[0].shapes
    table_shape = next((s for s in shapes if s.has_table), None)
    assert table_shape is not None
    assert table_shape.table.cell(0, 0).text == "Metric"
    assert table_shape.table.cell(1, 1).text == "$100M"


def test_chart_opens_in_powerpoint_shape(tmp_path):
    f = tmp_path / "d.pptx"
    make_deck(str(f), slides=1)

    data = {"cats": ["Q1", "Q2"], "series": {"Rev": [10, 20]}}
    res = add_chart(str(f), 1, "column", data)
    assert res["ok"] is True

    prs = Presentation(str(f))
    chart_shape = next((s for s in prs.slides[0].shapes if s.has_chart), None)
    assert chart_shape is not None


def test_add_image(tmp_path):
    f = tmp_path / "d.pptx"
    img_path = tmp_path / "test.png"
    make_deck(str(f), slides=1)
    make_test_image(str(img_path))

    res = add_image(str(f), 1, str(img_path))
    assert res["ok"] is True

    prs = Presentation(str(f))
    # Picture shape
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    pic_shape = next((s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE), None)
    assert pic_shape is not None


def test_preview_diff_never_writes(tmp_path):
    f = tmp_path / "d.pptx"
    make_deck(str(f), slides=1)

    with open(f, "rb") as fh:
        original_bytes = fh.read()

    diff = preview_diff(str(f), {"tool": "add_slide", "layout": "blank"})
    assert diff["ok"] is True
    assert "diff" in diff

    with open(f, "rb") as fh:
        after_bytes = fh.read()

    assert original_bytes == after_bytes
